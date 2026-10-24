"""Dòng thời gian xoắn ốc - Các nút sắp xếp theo đường xoắn ốc từ trong ra ngoài, mỗi nút là hình tròn màu kèm số thứ tự"""

import math
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from python.base import TemplateGenerator


class SpiralTimeline(TemplateGenerator):
    """Dòng thời gian xoắn ốc"""

    def __init__(self, theme=None):
        super().__init__(theme)
        self._title = ""
        self._subtitle = ""
        self._events = []  # list of dicts: {label, date, detail, color}
        self._node_radius = Inches(0.55)
        self._start_radius = Inches(0.3)
        self._default_colors = [
            RGBColor(52, 152, 219),
            RGBColor(231, 76, 60),
            RGBColor(46, 204, 113),
            RGBColor(155, 89, 182),
            RGBColor(241, 196, 15),
            RGBColor(230, 126, 34),
            RGBColor(26, 188, 156),
            RGBColor(243, 156, 18),
        ]

    def set_title(self, title, subtitle=""):
        self._title = title
        self._subtitle = subtitle
        return self

    def set_data(self, events):
        """Thiết lập danh sách sự kiện

        Args:
            events: list of dict, mỗi mục chứa:
                label (str): Tên sự kiện
                date (str): Mốc thời gian
                detail (str): Mô tả chi tiết
                color (RGBColor, optional): Màu tùy chỉnh
        """
        self._events = events
        return self

    def generate(self):
        """Sinh slide dòng thời gian xoắn ốc"""
        self.create_slide()
        slide_w = self.prs.slide_width
        slide_h = self.prs.slide_height

        # Tiêu đề
        if self._title:
            self.add_textbox(
                self._title, Inches(0.5), Inches(0.2),
                slide_w - Inches(1), Inches(0.5),
                font_size=26, bold=True,
                font_color=RGBColor(44, 62, 80),
                alignment=PP_ALIGN.LEFT
            )
        if self._subtitle:
            self.add_textbox(
                self._subtitle, Inches(0.5), Inches(0.65),
                slide_w - Inches(1), Inches(0.35),
                font_size=13,
                font_color=RGBColor(127, 140, 141),
                alignment=PP_ALIGN.LEFT
            )

        n = len(self._events)
        if n == 0:
            return self

        # Tâm xoắn ốc
        center_x = int(slide_w * 0.48)
        center_y = int(slide_h * 0.55)

        # Tham số xoắn ốc
        max_radius = min(slide_w, slide_h) * 0.38
        angle_step = 2 * math.pi / n  # Độ tăng góc giữa các nút
        radius_step = max_radius / max(n, 1)  # Bước tăng bán kính

        prev_x = None
        prev_y = None

        for i, event in enumerate(self._events):
            # Tọa độ xoắn ốc: góc tăng dần, bán kính tăng dần
            angle = angle_step * i - math.pi / 2  # Bắt đầu từ phía trên
            radius = self._start_radius + radius_step * (i + 1) * 0.5

            x = center_x + int(radius * math.cos(angle))
            y = center_y + int(radius * math.sin(angle))

            color = event.get("color") or self._default_colors[i % len(self._default_colors)]

            # Nối với nút trước đó
            if prev_x is not None:
                connector = self.slide.shapes.add_connector(
                    1, prev_x, prev_y, x, y
                )
                connector.line.color.rgb = RGBColor(189, 195, 199)
                connector.line.width = Pt(1.5)
                connector.line.dash_style = 2  # nét đứt

            prev_x = x
            prev_y = y

            # Nút tròn (bán kính giảm nhẹ theo vị trí)
            r = max(Inches(0.35), self._node_radius - Inches(0.03) * i)
            shape = self.add_shape(
                MSO_SHAPE.OVAL,
                x - r, y - r, r * 2, r * 2,
                fill_color=color
            )

            # Số thứ tự trong nút
            shape.text_frame.clear()
            p = shape.text_frame.paragraphs[0]
            run = p.add_run()
            run.text = str(i + 1)
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            p.alignment = PP_ALIGN.CENTER
            shape.text_frame.word_wrap = False

            # Nhãn văn bản cạnh nút
            label = event.get("label", "")
            date_str = event.get("date", "")

            # Quyết định hướng lệch văn bản dựa trên vị trí so với tâm
            dx = x - center_x
            dy = y - center_y
            if abs(dx) > abs(dy):
                # Lệch ngang lớn -> đặt văn bản ở trên hoặc dưới
                text_offset_x = -Inches(0.7)
                text_offset_y = r + Inches(0.1) if dy >= 0 else -r - Inches(0.45)
            else:
                # Lệch dọc lớn -> đặt văn bản ở trái hoặc phải
                text_offset_x = r + Inches(0.1) if dx >= 0 else -r - Inches(1.5)
                text_offset_y = -Inches(0.15)

            text_x = x + text_offset_x
            text_y = y + text_offset_y

            # Mốc thời gian
            self.add_textbox(
                date_str, text_x, text_y,
                Inches(1.3), Inches(0.2),
                font_size=9, bold=True,
                font_color=color,
                alignment=PP_ALIGN.CENTER
            )
            # Tên nhãn
            self.add_textbox(
                label, text_x, text_y + Inches(0.2),
                Inches(1.3), Inches(0.25),
                font_size=10, bold=True,
                font_color=RGBColor(44, 62, 80),
                alignment=PP_ALIGN.CENTER
            )

        # Điểm đánh dấu tâm
        self.add_shape(
            MSO_SHAPE.OVAL,
            center_x - Inches(0.12), center_y - Inches(0.12),
            Inches(0.24), Inches(0.24),
            fill_color=RGBColor(189, 195, 199)
        )

        return self


if __name__ == "__main__":
    timeline = SpiralTimeline()
    timeline.set_title("Lộ trình lặp sản phẩm", "Spiral Roadmap")
    timeline.set_data([
        {"label": "Kiểm chứng khái niệm", "date": "2024-01", "detail": "Phát triển MVP"},
        {"label": "Alpha", "date": "2024-03", "detail": "Thử nghiệm nội bộ"},
        {"label": "Beta", "date": "2024-05", "detail": "Thử nghiệm cộng đồng"},
        {"label": "v1.0", "date": "2024-07", "detail": "Phát hành chính thức"},
        {"label": "v1.5", "date": "2024-09", "detail": "Mở rộng tính năng"},
        {"label": "v2.0", "date": "2024-11", "detail": "Nâng cấp kiến trúc"},
        {"label": "Quốc tế hóa", "date": "2025-01", "detail": "Thị trường nước ngoài"},
        {"label": "Xây dựng hệ sinh thái", "date": "2025-03", "detail": "Nền tảng mở"},
    ])
    timeline.save("output/spiral_timeline.pptx")
    print("Đã sinh dòng thời gian xoắn ốc: output/spiral_timeline.pptx")
