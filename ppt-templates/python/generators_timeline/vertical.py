"""Dòng thời gian dọc - Trục thẳng đứng nằm bên trái, các sự kiện trải dài sang phải, chấm màu xen kẽ"""

from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from python.base import TemplateGenerator


class VerticalTimeline(TemplateGenerator):
    """Dòng thời gian dọc"""

    def __init__(self, theme=None):
        super().__init__(theme)
        self._title = ""
        self._subtitle = ""
        self._events = []  # list of dicts: {label, date, detail, color}
        self._node_radius = Inches(0.12)
        self._line_color = RGBColor(189, 195, 199)
        self._default_colors = [
            RGBColor(52, 152, 219),
            RGBColor(231, 76, 60),
            RGBColor(46, 204, 113),
            RGBColor(155, 89, 182),
            RGBColor(241, 196, 15),
            RGBColor(230, 126, 34),
        ]

    def set_title(self, title, subtitle=""):
        self._title = title
        self._subtitle = subtitle
        return self

    def set_data(self, events):
        """Thiết lập danh sách sự kiện

        Args:
            events: list of dict, mỗi mục chứa:
                label (str): Tên sự kiện
                date (str): Mốc thời gian
                detail (str): Mô tả chi tiết
                color (RGBColor, optional): Màu tùy chỉnh
        """
        self._events = events
        return self

    def generate(self):
        """Sinh slide dòng thời gian dọc"""
        self.create_slide()
        slide_w = self.prs.slide_width
        slide_h = self.prs.slide_height

        # Tiêu đề
        if self._title:
            self.add_textbox(
                self._title, Inches(0.5), Inches(0.2),
                slide_w - Inches(1), Inches(0.5),
                font_size=26, bold=True,
                font_color=RGBColor(44, 62, 80),
                alignment=PP_ALIGN.LEFT
            )
        if self._subtitle:
            self.add_textbox(
                self._subtitle, Inches(0.5), Inches(0.65),
                slide_w - Inches(1), Inches(0.35),
                font_size=13,
                font_color=RGBColor(127, 140, 141),
                alignment=PP_ALIGN.LEFT
            )

        # Trục dọc chính
        line_x = Inches(2.2)
        top_margin = Inches(1.2)
        bottom_margin = Inches(0.4)
        line_top = int(top_margin)
        line_bottom = int(slide_h - bottom_margin)

        # Trục chính
        axis_line = self.slide.shapes.add_connector(
            1, line_x, line_top, line_x, line_bottom
        )
        axis_line.line.color.rgb = self._line_color
        axis_line.line.width = Pt(3)

        n = len(self._events)
        if n == 0:
            return self

        available_height = line_bottom - line_top
        spacing = available_height / max(n, 1)

        for i, event in enumerate(self._events):
            y = line_top + int(spacing * (i + 0.5))
            color = event.get("color") or self._default_colors[i % len(self._default_colors)]

            # Nút tròn
            r = self._node_radius
            self.add_shape(
                MSO_SHAPE.OVAL,
                line_x - r, y - r, r * 2, r * 2,
                fill_color=color
            )

            # Đường nối ngang
            connector = self.slide.shapes.add_connector(
                1, line_x + r + Inches(0.05), y,
                line_x + Inches(1.0), y
            )
            connector.line.color.rgb = color
            connector.line.width = Pt(1.5)

            # Nhãn mốc thời gian - Nằm bên trái trục dọc
            date_str = event.get("date", "")
            self.add_textbox(
                date_str,
                Inches(0.3), y - Inches(0.15),
                Inches(1.7), Inches(0.3),
                font_size=10, bold=True,
                font_color=color,
                alignment=PP_ALIGN.RIGHT
            )

            # Tiêu đề và mô tả sự kiện - Nằm bên phải trục dọc
            label = event.get("label", "")
            detail = event.get("detail", "")

            text_x = line_x + Inches(1.1)
            text_w = slide_w - text_x - Inches(0.3)

            # Tiêu đề
            self.add_textbox(
                label,
                text_x, y - Inches(0.25),
                text_w, Inches(0.3),
                font_size=13, bold=True,
                font_color=RGBColor(44, 62, 80),
                alignment=PP_ALIGN.LEFT
            )
            # Mô tả
            if detail:
                self.add_textbox(
                    detail,
                    text_x, y + Inches(0.05),
                    text_w, Inches(0.35),
                    font_size=10,
                    font_color=RGBColor(127, 140, 141),
                    alignment=PP_ALIGN.LEFT
                )

        return self


if __name__ == "__main__":
    timeline = VerticalTimeline()
    timeline.set_title("Lịch sử phát triển công ty", "Company History")
    timeline.set_data([
        {"label": "Thành lập công ty", "date": "2020-01", "detail": "Đăng ký thành lập, quy tụ đội ngũ nòng cốt"},
        {"label": "Vòng vốn thiên thần", "date": "2020-06", "detail": "Nhận 5 triệu USD vốn đầu tư thiên thần"},
        {"label": "Ra mắt sản phẩm", "date": "2021-03", "detail": "Phát hành phiên bản v1.0 chính thức"},
        {"label": "Gọi vốn Vòng A", "date": "2022-01", "detail": "Hoàn tất gọi vốn Vòng A trị giá 30 triệu USD"},
        {"label": "Mở rộng quy mô", "date": "2022-09", "detail": "Quy mô nhân sự vượt mốc 100 thành viên"},
        {"label": "Vươn ra toàn cầu", "date": "2023-06", "detail": "Tiến quân vào thị trường quốc tế"},
        {"label": "Gọi vốn Vòng B", "date": "2024-01", "detail": "Hoàn tất gọi vốn Vòng B"},
    ])
    timeline.save("output/vertical_timeline.pptx")
    print("Đã sinh dòng thời gian dọc: output/vertical_timeline.pptx")
