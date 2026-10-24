"""Bộ sinh trang bìa chia đôi màn hình - Thiết kế phân chia trái phải, tương phản giữa màu sắc và khoảng trắng"""

from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from python.base import TemplateGenerator
from python.utils import hex_to_rgb


class SplitScreenCover(TemplateGenerator):
    """Bìa chia đôi màn hình"""

    def __init__(self, theme=None):
        super().__init__(theme)
        self._title = ""
        self._subtitle = ""
        self._title_font_size = 40
        self._subtitle_font_size = 18
        self._left_color = None
        self._right_color = None
        self._title_color = RGBColor(255, 255, 255)
        self._subtitle_color = RGBColor(60, 60, 60)
        self._split_ratio = 0.5  # 0.0~1.0 Tỷ lệ phần bên trái

    def set_title(self, title, font_size=None, color=None):
        """Thiết lập tiêu đề chính (hiển thị ở mảng màu bên trái)"""
        self._title = title
        if font_size is not None:
            self._title_font_size = font_size
        if color is not None:
            self._title_color = color
        return self

    def set_subtitle(self, subtitle, font_size=None, color=None):
        """Thiết lập tiêu đề phụ (hiển thị ở mảng sáng bên phải)"""
        self._subtitle = subtitle
        if font_size is not None:
            self._subtitle_font_size = font_size
        if color is not None:
            self._subtitle_color = color
        return self

    def set_split_ratio(self, ratio):
        """Thiết lập tỷ lệ phân chia màn hình (0.0~1.0), tỷ lệ phần bên trái"""
        self._split_ratio = max(0.1, min(0.9, ratio))
        return self

    def _get_theme_colors(self):
        """Lấy cấu hình màu chủ đề"""
        if self.theme:
            primary = hex_to_rgb(self.theme.primary)
            secondary = hex_to_rgb(self.theme.secondary)
            accent = hex_to_rgb(self.theme.accent)
        else:
            primary = RGBColor(44, 62, 80)
            secondary = RGBColor(52, 152, 219)
            accent = RGBColor(231, 76, 60)
        return primary, secondary, accent

    def generate(self):
        """Sinh slide trang bìa chia đôi màn hình"""
        self.create_slide()
        slide_w = self.prs.slide_width
        slide_h = self.prs.slide_height
        primary, secondary, accent = self._get_theme_colors()

        if self._left_color is None:
            self._left_color = primary
        if self._right_color is None:
            self._right_color = RGBColor(250, 250, 250)

        split_x = int(slide_w * self._split_ratio)
        right_w = slide_w - split_x

        # --- Bên trái: Khối màu gradient ---
        left_bg = self.add_gradient_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0), Inches(0), split_x, slide_h,
            [primary, secondary], angle=135
        )

        # --- Bên phải: Nền màu sáng ---
        right_bg = self.add_shape(
            MSO_SHAPE.RECTANGLE,
            split_x, Inches(0), right_w, slide_h,
            fill_color=self._right_color
        )

        # --- Đường phân cách dọc ---
        line_color = accent
        self.add_shape(
            MSO_SHAPE.RECTANGLE,
            split_x - Pt(1), Inches(0),
            Pt(3), slide_h,
            fill_color=line_color
        )

        # --- Họa tiết hình học bên trái ---
        deco_size = Inches(1.8)
        self.add_shape(
            MSO_SHAPE.OVAL,
            split_x - deco_size - Inches(0.3),
            Inches(0.5),
            deco_size, deco_size,
            fill_color=accent
        )
        from python.utils import set_shape_transparency
        set_shape_transparency(self.slide.shapes[-1], 0.75)

        # --- Tiêu đề bên trái ---
        margin = Inches(0.8)
        self.add_textbox(
            self._title,
            margin, slide_h // 2 - Inches(0.6),
            split_x - margin * 2, Inches(1.2),
            font_size=self._title_font_size,
            font_color=self._title_color,
            bold=True,
            alignment=PP_ALIGN.LEFT
        )

        # --- Tiêu đề phụ bên phải ---
        self.add_textbox(
            self._subtitle,
            split_x + Inches(0.6),
            slide_h // 2 - Inches(0.4),
            right_w - Inches(1.2), Inches(0.8),
            font_size=self._subtitle_font_size,
            font_color=self._subtitle_color,
            bold=False,
            alignment=PP_ALIGN.LEFT
        )

        # --- Ô vuông nhỏ trang trí góc dưới bên phải ---
        block_size = Inches(0.5)
        self.add_shape(
            MSO_SHAPE.RECTANGLE,
            split_x + right_w - block_size - Inches(0.6),
            slide_h - block_size - Inches(0.6),
            block_size, block_size,
            fill_color=secondary
        )
        set_shape_transparency(self.slide.shapes[-1], 0.6)

        return self


if __name__ == "__main__":
    cover = SplitScreenCover()
    cover.set_title("Bìa chia đôi màn hình", font_size=44)
    cover.set_subtitle("Split Screen Cover", font_size=20)
    cover.save("output/split_screen_cover.pptx")
    print("Đã sinh bìa chia đôi màn hình: output/split_screen_cover.pptx")
