#!/usr/bin/env python3
"""Xác minh các mẫu PPT sinh ra có chứa nội dung thực chất hay không.

Các mẫu của dự án này được sinh hàng loạt bằng script; nếu một bộ sinh nào đó không nhận được dữ liệu mẫu,
sản phẩm tạo ra có thể chỉ có tiêu đề (hoặc thậm chí không có slide nào), trong khi quá trình sinh không hề báo lỗi.
Script này giúp phát hiện rõ tình trạng "sinh thành công nhưng rỗng nội dung" để CI ngăn chặn kịp thời.

Cách dùng:
    python scripts/verify_templates.py <thư_mục>
    python scripts/verify_templates.py templates --min-shapes 4

Quy tắc phán đoán:
    - 0 trang: Số slide bằng 0
    - Chỉ có tiêu đề: Số slide ≥ 1, nhưng số hình khối ở trang đầu < --min-shapes

Mã thoát 0 nghĩa là vượt qua tất cả, 1 nghĩa là có sản phẩm không đạt tiêu chuẩn.
"""

from __future__ import annotations

import argparse
import sys
from collections import Counter
from pathlib import Path

from pptx import Presentation

DEFAULT_MIN_SHAPES = 4


def inspect(path: Path) -> tuple[int, int]:
    """Trả về (số trang slide, số hình khối ở trang đầu)"""
    prs = Presentation(str(path))
    page_count = len(prs.slides)
    shape_count = len(prs.slides[0].shapes) if page_count else 0
    return page_count, shape_count


def main() -> int:
    parser = argparse.ArgumentParser(description="Xác minh mẫu pptx có chứa nội dung thực chất hay không")
    parser.add_argument("directory", help="Thư mục mẫu (sẽ tìm đệ quy *.pptx)")
    parser.add_argument(
        "--min-shapes",
        type=int,
        default=DEFAULT_MIN_SHAPES,
        help=f"Ngưỡng tối thiểu số hình khối trang đầu, dưới ngưỡng này coi là trang chỉ có tiêu đề (mặc định {DEFAULT_MIN_SHAPES})",
    )
    parser.add_argument("-v", "--verbose", action="store_true", help="Liệt kê chi tiết từng file không đạt")
    args = parser.parse_args()

    root = Path(args.directory)
    if not root.is_dir():
        print(f"[FAIL] Thư mục không tồn tại: {root}", file=sys.stderr)
        return 1

    files = sorted(root.rglob("*.pptx"))
    if not files:
        print(f"[FAIL] Không tìm thấy file .pptx nào trong thư mục: {root}", file=sys.stderr)
        return 1

    empty: list[str] = []      # 0 trang
    title_only: list[str] = []  # Chỉ có tiêu đề
    groups: Counter[str] = Counter()
    shape_stats: list[int] = []

    for path in files:
        rel = path.relative_to(root).as_posix()
        pages, shapes = inspect(path)
        groups[path.parent.relative_to(root).as_posix()] += 1

        if pages == 0:
            empty.append(rel)
        elif shapes < args.min_shapes:
            title_only.append(f"{rel} (trang đầu có {shapes} hình khối)")
        else:
            shape_stats.append(shapes)

    print(f"Quét thư mục: {root}")
    print(f"Tổng số mẫu: {len(files)}    Ngưỡng hình khối tối thiểu trang đầu: {args.min_shapes}")
    print("-" * 60)
    for group, count in sorted(groups.items()):
        print(f"  {group:<28}{count:>6}")
    print("-" * 60)

    failed = len(empty) + len(title_only)
    print(f"Mẫu 0 trang           : {len(empty)}")
    print(f"Mẫu chỉ có tiêu đề    : {len(title_only)}")
    if shape_stats:
        print(f"Mẫu đạt tiêu chuẩn    : {len(shape_stats)} (Trung vị số hình khối trang đầu {sorted(shape_stats)[len(shape_stats) // 2]})")

    if args.verbose and (empty or title_only):
        print("\nChi tiết các mẫu không đạt:")
        for rel in empty:
            print(f"  [0 trang]          {rel}")
        for rel in title_only:
            print(f"  [Chỉ có tiêu đề]   {rel}")

    if failed:
        print(f"\n[FAIL] {failed} mẫu thiếu nội dung thực chất (trong tổng số {len(files)} mẫu).", file=sys.stderr)
        return 1

    print(f"\n[OK] Toàn bộ {len(files)} mẫu đều chứa nội dung thực chất.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
