#!/usr/bin/env python3
"""
Bộ tạo Báo cáo Chuyên đề / Bài tập lớn môn Quản lý Quy trình Nghiệp vụ (BPM)
Học viện Ngân hàng - Khoa Sau Đại học
Tác giả: Đỗ Văn Đạt - 28k405001 - Lớp K28PTKDA
Đề tài: TỐI ƯU HÓA QUY TRÌNH TRIỂN KHAI TÍN DỤNG VÀ GIẢI NGÂN VỐN VAY TẠI LPBANK THEO CHUẨN APQC PCF VÀ KIẾN TRÚC SMART BPM
"""

import os
import sys
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

BASE_DIR = Path('/Users/dovandat/WORKSPACE/HVNH/BPM/Bai_tap_lon')
TEMPLATE_PATH = Path('/Users/dovandat/.gemini/config/skills/academic-docx-writer/templates/DoVanDat_CSDLNC_template.docx')
OUTPUT_DOCX = BASE_DIR / 'Bao_cao_BPM_Tin_dung_Giai_ngan_LPBank_HOAN_THIEN.docx'
OUTPUT_MD = BASE_DIR / 'Bao_cao_BPM_Tin_dung_Giai_ngan_LPBank_HOAN_THIEN.md'
DIAGRAM_DIR = BASE_DIR / 'diagrams'

# Màu sắc chuẩn học thuật HVNH
COLOR_NAVY_HVNH   = RGBColor(10, 37, 64)       # #0A2540
COLOR_BLUE_SEC    = RGBColor(31, 77, 120)      # #1F4D78
COLOR_ACCENT_BLUE = RGBColor(46, 116, 181)     # #2E74B5
COLOR_DARK_TEXT   = RGBColor(15, 23, 42)       # #0F172A
COLOR_MUTED_TEXT  = RGBColor(100, 116, 139)    # #64748B

HEX_HDR_BG   = "D9E2F3"
HEX_ZEBRA_BG = "F8FAFC"
HEX_BORDER   = "94A3B8"
HEX_CALLOUT_BG = "F1F5F9"

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=60, bottom=60, left=110, right=110):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, border_color=HEX_BORDER):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>'
        f'<w:left w:val="none"/>'
        f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="{border_color}"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def make_tbl_header(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

def make_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

def add_body_p(doc, text, bold_prefix=None, indent=True, line_spacing=1.5, space_after=6):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.first_line_indent = Inches(0.5) if indent else Inches(0)
        
    if bold_prefix:
        r_b = p.add_run(bold_prefix)
        r_b.font.name = 'Times New Roman'
        r_b.font.size = Pt(13)
        r_b.font.bold = True
        r_b.font.color.rgb = COLOR_DARK_TEXT
        
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(13)
    r.font.color.rgb = COLOR_DARK_TEXT
    return p

def add_bullet_p(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.35
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.first_line_indent = Inches(0)
    
    if bold_prefix:
        r_b = p.add_run(bold_prefix)
        r_b.font.name = 'Times New Roman'
        r_b.font.size = Pt(13)
        r_b.font.bold = True
        r_b.font.color.rgb = COLOR_DARK_TEXT
        
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(13)
    r.font.color.rgb = COLOR_DARK_TEXT
    return p

def add_h1(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.first_line_indent = Inches(0)
    
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(15.5)
    r.font.bold = True
    r.font.color.rgb = COLOR_NAVY_HVNH
    return p

def add_h2(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.first_line_indent = Inches(0)
    
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(13.5)
    r.font.bold = True
    r.font.color.rgb = COLOR_BLUE_SEC
    return p

def add_h3(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.first_line_indent = Inches(0)
    
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = COLOR_ACCENT_BLUE
    return p

def add_callout(doc, text, bold_title="Chú giải chuyên sâu:"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, HEX_CALLOUT_BG)
    set_cell_margins(cell, top=100, bottom=100, left=160, right=160)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="0A2540"/>'
        f'<w:top w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.space_after = Pt(0)
    
    if bold_title:
        r_t = p.add_run(bold_title + " ")
        r_t.font.name = 'Times New Roman'
        r_t.font.size = Pt(11.5)
        r_t.font.bold = True
        r_t.font.color.rgb = COLOR_NAVY_HVNH
        
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11.5)
    r.font.color.rgb = COLOR_DARK_TEXT
    
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(4)
    p_after.paragraph_format.line_spacing = 1.0

def add_fig(doc, img_path, caption_text, width_in=6.6):
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(8)
    p_img.paragraph_format.space_after = Pt(4)
    r_img = p_img.add_run()
    r_img.add_picture(str(img_path), width=Inches(width_in))
    
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(10)
    p_cap.paragraph_format.keep_with_next = False
    
    r_cap = p_cap.add_run(caption_text)
    r_cap.font.name = 'Times New Roman'
    r_cap.font.size = Pt(11)
    r_cap.font.italic = True
    r_cap.font.color.rgb = COLOR_MUTED_TEXT

def add_tbl(doc, headers, data, caption_text, col_widths=None):
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_cap.paragraph_format.space_before = Pt(10)
    p_cap.paragraph_format.space_after = Pt(4)
    p_cap.paragraph_format.keep_with_next = True
    
    r_cap = p_cap.add_run(caption_text)
    r_cap.font.name = 'Times New Roman'
    r_cap.font.size = Pt(11.5)
    r_cap.font.bold = True
    r_cap.font.color.rgb = COLOR_NAVY_HVNH
    
    rows = len(data) + 1
    cols = len(headers)
    table = doc.add_table(rows=rows, cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    
    # Header
    hdr_row = table.rows[0]
    make_tbl_header(hdr_row)
    make_cant_split(hdr_row)
    for j, h_text in enumerate(headers):
        cell = hdr_row.cells[j]
        set_cell_background(cell, HEX_HDR_BG)
        set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(h_text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(11)
        r.font.bold = True
        r.font.color.rgb = COLOR_NAVY_HVNH
        
    # Data rows
    for i, row_data in enumerate(data):
        row = table.rows[i + 1]
        make_cant_split(row)
        bg = HEX_ZEBRA_BG if i % 2 == 1 else "FFFFFF"
        for j, val in enumerate(row_data):
            cell = row.cells[j]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j > 1 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(str(val))
            r.font.name = 'Times New Roman'
            r.font.size = Pt(10.5)
            r.font.color.rgb = COLOR_DARK_TEXT
            
    if col_widths and len(col_widths) == cols:
        for row in table.rows:
            for j, w in enumerate(col_widths):
                row.cells[j].width = Inches(w)
                
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(6)
    p_after.paragraph_format.line_spacing = 1.0

print(">>> Cấu hình cơ bản đã hoàn tất.")
