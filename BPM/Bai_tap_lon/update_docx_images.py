#!/usr/bin/env python3
"""
Script cập nhật 4 hình ảnh Retina 2x mới vào file Word (.docx)
Bảo đảm tính thẩm mỹ, độ phân giải cao và chuẩn mực định dạng văn bản học thuật HVNH.
"""

import os
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

BASE_DIR = Path('/Users/dovandat/WORKSPACE/HVNH/BPM/Bai_tap_lon')
DOCX_PATH = BASE_DIR / 'Bao_cao_BPM_Tin_dung_Giai_ngan_LPBank_HOAN_THIEN.docx'
DIAGRAM_DIR = BASE_DIR / 'diagrams'

IMAGES_MAP = [
    {
        'key': 'Hình 1',
        'img_path': DIAGRAM_DIR / 'Hinh1_Ban_do_chien_luoc_LPBank.drawio.png',
        'new_caption': 'Hình 1. Bản đồ chiến lược cấp quy trình giải ngân vốn vay theo mô hình Balanced Scorecard (LPBank). Nguồn: Nhóm tác giả xây dựng.',
    },
    {
        'key': 'Hình 2',
        'img_path': DIAGRAM_DIR / 'Hinh2_Process_Landscape_LPBank.drawio.png',
        'new_caption': 'Hình 2. Sơ đồ cảnh quan quy trình cấp cao (Process Landscape) và Chuỗi giá trị tín dụng LPBank theo chuẩn APQC PCF Banking v7.2.2. Nguồn: Nhóm tác giả tổng hợp.',
    },
    {
        'key': 'Hình 3',
        'img_path': DIAGRAM_DIR / 'Hinh3_Quy_trinh_Hien_trang_AsIs_LPBank.drawio.png',
        'new_caption': 'Hình 3. Sơ đồ luồng quy trình hiện trạng (As-Is) triển khai giải ngân vốn vay LPBank chuẩn BPMN 2.0 (Minh họa 3 điểm nghẽn Bottlenecks và 2 vòng lặp Rework). Nguồn: Nhóm tác giả mô hình hóa từ hồ sơ thiết kế [6].',
    },
    {
        'key': 'Hình 4',
        'img_path': DIAGRAM_DIR / 'Hinh4_Quy_trinh_Tuong_lai_ToBe_LPBank.drawio.png',
        'new_caption': 'Hình 4. Sơ đồ luồng quy trình tương lai (To-Be) triển khai giải ngân vốn vay LPBank chuẩn BPMN 2.0 (Kiến trúc Smart BPM Orchestrator, Ký số SmartCA và Xử lý giao dịch thẳng STP tự động 100%). Nguồn: Nhóm tác giả đề xuất.',
    },
]

def update_docx():
    print(f"Đang đọc file Word: {DOCX_PATH}")
    doc = docx.Document(str(DOCX_PATH))
    
    for item in IMAGES_MAP:
        key = item['key']
        img_path = str(item['img_path'])
        caption = item['new_caption']
        
        # Tìm paragraph chứa key caption
        caption_idx = -1
        for i, p in enumerate(doc.paragraphs):
            if key in p.text:
                caption_idx = i
                break
                
        if caption_idx == -1:
            print(f"[!] Không tìm thấy caption cho {key}")
            continue
            
        print(f"[+] Tìm thấy {key} tại paragraph {caption_idx}")
        
        # Cập nhật caption text
        cap_p = doc.paragraphs[caption_idx]
        cap_p.text = caption
        cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in cap_p.runs:
            run.font.name = 'Times New Roman'
            run.font.size = Pt(10.5)
            run.font.italic = True
            run.font.color.rgb = RGBColor(100, 116, 139) # slate-500
            
        # Paragraph trước đó chứa ảnh
        img_p = doc.paragraphs[caption_idx - 1]
        # Xóa nội dung / hình ảnh cũ trong img_p
        p_elem = img_p._p
        p_elem.clear_content()
        img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # Chèn ảnh mới độ nét cao Retina 2x
        run = img_p.add_run()
        run.add_picture(img_path, width=Inches(6.6))
        print(f"    [✓] Đã chèn {img_path} (width=6.6 inches) vào paragraph {caption_idx - 1}")

    doc.save(str(DOCX_PATH))
    print(f">>> HOÀN TẤT CẬP NHẬT FILE WORD: {DOCX_PATH}")

if __name__ == '__main__':
    update_docx()
