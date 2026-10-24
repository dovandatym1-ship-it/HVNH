import docx
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Pt, Inches

doc = docx.Document()

# Tạo table mẫu
table = doc.add_table(rows=2, cols=2)
# Thêm border chuẩn
tblPr = table._tbl.tblPr
borders = parse_xml(
    f'<w:tblBorders {nsdecls("w")}>'
    f'<w:top w:val="single" w:sz="6" w:space="0" w:color="94A3B8"/>'
    f'<w:left w:val="none"/>'
    f'<w:bottom w:val="single" w:sz="8" w:space="0" w:color="94A3B8"/>'
    f'<w:right w:val="none"/>'
    f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="94A3B8"/>'
    f'<w:insideV w:val="none"/>'
    f'</w:tblBorders>'
)
# Tìm tblLook trong tblPr để chèn borders TRƯỚC tblLook
tblLook = tblPr.find(docx.oxml.ns.qn('w:tblLook'))
if tblLook is not None:
    tblLook.addprevious(borders)
else:
    tblPr.append(borders)

for r_idx, row in enumerate(table.rows):
    for c_idx, cell in enumerate(row.cells):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="D9E2F3"/>')
        tcPr.append(shd)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        p.text = f"Ô {r_idx},{c_idx}"

doc.save('/tmp/test_perfect_schema.docx')
