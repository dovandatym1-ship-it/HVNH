import docx

TEMPLATE_PATH = '/Users/dovandat/.gemini/config/skills/academic-docx-writer/templates/DoVanDat_CSDLNC_template.docx'
doc = docx.Document(TEMPLATE_PATH)
body = doc._element.body

print("Ban đầu: số elements =", len(body))

# Tìm vị trí của Table 1 (Table từ viết tắt)
table_count = 0
tbl_1_idx = -1
for i, el in enumerate(body):
    if el.tag.endswith('tbl'):
        table_count += 1
        if table_count == 2: # Table 1 (0-indexed là table thứ 2)
            tbl_1_idx = i
            break

print(f"Table 1 tìm thấy tại index: {tbl_1_idx}")

# Thu thập các elements cần xóa từ tbl_1_idx + 1 đến hết (trừ sectPr cuối cùng)
to_remove = []
for i in range(tbl_1_idx + 1, len(body)):
    el = body[i]
    if not el.tag.endswith('sectPr'):
        to_remove.append(el)

print(f"Số phần tử cần xóa sạch: {len(to_remove)}")
for el in to_remove:
    body.remove(el)

print("Sau khi dọn sạch: số elements =", len(body))
print("Các tables còn lại:", len(doc.tables))
for i, t in enumerate(doc.tables):
    print(f"Table {i}: rows={len(t.rows)}, cols={len(t.columns)}")
