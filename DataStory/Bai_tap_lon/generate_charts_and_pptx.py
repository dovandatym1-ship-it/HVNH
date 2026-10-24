import os
import sys
import subprocess

try:
    import pandas as pd
    import numpy as np
    import matplotlib.pyplot as plt
    import seaborn as sns
    from pptx import Presentation
    from pptx.util import Inches, Pt
except ImportError:
    subprocess.check_call([sys.executable, "-m", "pip", "install", "pandas", "matplotlib", "seaborn", "python-pptx"])
    import pandas as pd
    import numpy as np
    import matplotlib.pyplot as plt
    import seaborn as sns
    from pptx import Presentation
    from pptx.util import Inches, Pt

data_path = "DS_Cuoiky/cleaned_data.csv"
donut_chart_path = "DS_Cuoiky/donut_chart.png"
bar_chart_path = "DS_Cuoiky/bar_chart.png"
scatter_plot_path = "DS_Cuoiky/scatter_plot.png"
histogram_path = "DS_Cuoiky/histogram.png"
pptx_path = "DS_Cuoiky/Slide_Thuyet_Trinh_Chi_Tiet.pptx"

# Đọc dữ liệu
df_clean = pd.read_csv(data_path)
df_clean['Trạng thái'] = df_clean['Số ngày chậm'].apply(lambda x: 'Trễ hạn' if x > 0 else 'Đúng/Sớm hạn')
df_delayed = df_clean[df_clean['Số ngày chậm'] > 0]

sns.set_theme(style="whitegrid")

# --- TẠO BIỂU ĐỒ 1: DONUT CHART ---
plt.figure(figsize=(8, 6))
status_counts = df_clean['Trạng thái'].value_counts()
colors = ['seagreen', 'firebrick'] if status_counts.index[0] != 'Trễ hạn' else ['firebrick', 'seagreen']
plt.pie(status_counts, labels=status_counts.index, autopct='%1.1f%%', startangle=90, colors=colors, textprops={'fontsize': 14, 'weight': 'bold'})
centre_circle = plt.Circle((0,0),0.70,fc='white')
fig = plt.gcf()
fig.gca().add_artist(centre_circle)
plt.title('Tỷ lệ Dự án Trễ hạn (2022-2025)', fontsize=16, fontweight='bold', pad=20)
plt.tight_layout()
plt.savefig(donut_chart_path, dpi=150)
plt.close()

# --- TẠO BIỂU ĐỒ 2: BAR CHART ---
cause_impact = df_delayed.groupby('Nguyên nhân chậm tiến độ')['Số ngày chậm'].mean().sort_values(ascending=False).reset_index()
plt.figure(figsize=(10, 6))
ax = sns.barplot(data=cause_impact, x='Số ngày chậm', y='Nguyên nhân chậm tiến độ', hue='Nguyên nhân chậm tiến độ', palette='viridis', legend=False)
plt.title('Thời gian trễ hạn trung bình theo từng nguyên nhân', fontsize=16, fontweight='bold', pad=20)
plt.xlabel('Thời gian trễ hạn trung bình (Ngày)', fontsize=12)
plt.ylabel('')
for i, patch in enumerate(ax.patches):
    if cause_impact.iloc[i]['Nguyên nhân chậm tiến độ'] == 'Thay đổi thiết kế':
        patch.set_facecolor('firebrick')
plt.tight_layout()
plt.savefig(bar_chart_path, dpi=150)
plt.close()

# --- TẠO BIỂU ĐỒ 3: SCATTER PLOT ---
plt.figure(figsize=(12, 6))
sns.scatterplot(data=df_clean, x='Số lần thay đổi thiết kế', y='Số ngày chậm', hue='Trạng thái', 
                palette={'Trễ hạn': 'firebrick', 'Đúng/Sớm hạn': 'lightgrey'}, alpha=0.6, s=60)
plt.title('Tương quan Tỷ lệ thuận giữa Đổi thiết kế và Tiến độ', fontsize=16, fontweight='bold', pad=20)
plt.xlabel('Số lần thay đổi bản vẽ thiết kế', fontsize=12)
plt.ylabel('Số ngày chậm', fontsize=12)
df_trend = df_clean[df_clean['Trạng thái'] == 'Trễ hạn']
z = np.polyfit(df_trend['Số lần thay đổi thiết kế'], df_trend['Số ngày chậm'], 1)
p = np.poly1d(z)
plt.plot(df_trend['Số lần thay đổi thiết kế'], p(df_trend['Số lần thay đổi thiết kế']), "r--", linewidth=2)
plt.text(12, 50, 'Càng sửa bản vẽ, Dự án càng trượt dài!', fontsize=14, color='firebrick', fontweight='bold', bbox=dict(facecolor='white', alpha=0.8, edgecolor='none'))
plt.tight_layout()
plt.savefig(scatter_plot_path, dpi=150)
plt.close()

# --- TẠO BIỂU ĐỒ 4: HISTOGRAM ---
plt.figure(figsize=(10, 6))
sns.histplot(data=df_delayed, x='Số ngày chậm', bins=40, kde=True, color='firebrick')
plt.title('Phân phối Số ngày chậm (The Snowball Effect)', fontsize=16, fontweight='bold', pad=20)
plt.xlabel('Số ngày chậm (Ngày)', fontsize=12)
plt.ylabel('Số lượng dự án', fontsize=12)
avg_delay = df_delayed['Số ngày chậm'].mean()
plt.axvline(avg_delay, color='black', linestyle='--', linewidth=2)
plt.text(avg_delay + 5, plt.ylim()[1]*0.8, f'Trung bình: {avg_delay:.0f} ngày', color='black', fontsize=12, weight='bold')
plt.text(5, plt.ylim()[1]*0.5, 'Vùng trống\n(Thiếu khả năng phục hồi)', color='grey', fontsize=12, style='italic', ha='center')
plt.tight_layout()
plt.savefig(histogram_path, dpi=150)
plt.close()

print("Đã tạo và lưu các biểu đồ mới.")

# --- TẠO FILE PPTX VỚI FULL HÌNH ẢNH ---
slides_data = [
    {
        "type": "title",
        "title": "Không sợ 'Lỗ hổng' phát sinh, chỉ sợ thiếu 'Khả năng phục hồi'",
        "subtitle": "Giải pháp tối ưu tiến độ dự án xây dựng\n(Phân tích từ 6.000 dự án - Công ty CP Tập đoàn Công nghiệp Việt)",
        "image": "DS_Cuoiky/blueprint_hook_1783749990870.jpg",
        "notes": "Chào thầy và các bạn..."
    },
    {
        "type": "content_with_image",
        "title": "1. Bối cảnh (The Context)",
        "body": "BỨC TRANH TOÀN CẢNH:\n• Giai đoạn 2022-2025: Triển khai 6.000 dự án.\n• Thực trạng: Có tới 40% dự án bị trễ hạn.\n• Lỗ hổng khổng lồ nào đang nuốt chửng các dự án này?",
        "image": donut_chart_path,
        "notes": "Hãy cùng bước vào câu chuyện của Tập đoàn Công nghiệp Việt..."
    },
    {
        "type": "content_with_image",
        "title": "2. Nút thắt (The Conflict): Xóa bỏ định kiến",
        "body": "SỰ THẬT TỪ DỮ LIỆU:\n• Thời tiết, Mặt bằng chỉ đóng góp dưới 15%.\n• Kẻ kéo lùi dự án tàn bạo nhất lộ diện: THAY ĐỔI THIẾT KẾ!",
        "image": bar_chart_path,
        "notes": "Khi ban giám đốc đổ lỗi cho ông trời, dữ liệu Bar Chart đã bác bỏ hoàn toàn..."
    },
    {
        "type": "content_with_image",
        "title": "3. Điểm bùng nổ (The Aha! Moment)",
        "body": "SỰ TƯƠNG QUAN CHẾT CHÓC:\n• Càng sửa bản vẽ, dự án càng trượt dài!\n• Lỗ hổng không đến từ ngoại cảnh, mà đến từ quyết định của con người.",
        "image": scatter_plot_path,
        "notes": "Và đây chính là khoảnh khắc 'Aha!' – điểm sáng nhất trong toàn bộ bài phân tích..."
    },
    {
        "type": "content_with_image",
        "title": "4. Phân tích sâu (Deep Dive): Quả cầu tuyết",
        "body": "HIỆU ỨNG HÒN TUYẾT LĂN:\n• Vùng trễ hạn 1-10 ngày gần như trống.\n• Sự thiếu hụt KHẢ NĂNG PHỤC HỒI TIẾN ĐỘ khiến mọi rủi ro đều tích tụ thành thảm họa > 50 ngày.",
        "image": histogram_path,
        "notes": "Biểu đồ phân phối này chỉ ra sự thiếu hụt khả năng phục hồi..."
    },
    {
        "type": "content_with_image",
        "title": "5. Giải pháp Kép: Vá Lỗ hổng (Phần 1)",
        "body": "VÁ LỖ HỔNG TỪ PHÒNG THIẾT KẾ\n\n• Kỷ luật 'Design Freeze' (Khóa bản vẽ):\n  - Thiết lập một mốc thời gian tử thần trong hợp đồng.\n  - Mọi thay đổi sau mốc này bị cấm hoặc phải bị kiểm soát chặt chẽ.\n• Đừng chạy đua trên bùn lầy, hãy chốt chặn từ phòng điều hòa!",
        "image": "DS_Cuoiky/design_freeze_1783750009687.jpg",
        "notes": "Để giải cứu Tập đoàn, chúng em đề xuất 2 bước hành động. Bước 1: Vá lỗ hổng..."
    },
    {
        "type": "content_with_image",
        "title": "6. Giải pháp Kép: Tăng Phục hồi (Phần 2)",
        "body": "TĂNG CƯỜNG KHẢ NĂNG PHỤC HỒI (SCHEDULE RESILIENCE)\n\n• Quy trình 'Change Order' (Biên lai thay đổi):\n  - Yêu cầu đổi thiết kế = Tiền phạt + Thời gian gia hạn.\n• Kích hoạt nén tiến độ:\n  - Bơm ngân sách làm song song (Fast-tracking).\n  - Dập tắt độ trễ trong vòng 48h để chặn đứng quả cầu tuyết.",
        "image": "DS_Cuoiky/schedule_crashing_1783750017966.jpg",
        "notes": "Nhưng nếu khách hàng vẫn kiên quyết bắt ép đổi thiết kế thì sao? Đó là lúc ta cần Bước 2..."
    },
    {
        "type": "content_with_image",
        "title": "7. Hỗ trợ Công nghệ: Áp dụng mô hình BIM 3D",
        "body": "Tương lai của quản lý thiết kế:\n\n• Chuyển dịch từ bản vẽ 2D sang mô hình 3D (BIM).\n• Trực quan hóa công trình trước khi khởi công.\n• Giảm thiểu tối đa sự dao động tâm lý của khách hàng, hạn chế thay đổi thiết kế về sau.",
        "image": "DS_Cuoiky/bim_3d_1783750026895.jpg",
        "notes": "Ngoài pháp lý hợp đồng, việc chuyển dịch sang mô hình 3D (BIM) cũng là một vũ khí..."
    },
    {
        "type": "content_with_image",
        "title": "8. Lời kêu gọi (Call to Action)",
        "body": "Quản trị dự án xuất sắc không phải là người giăng lưới bắt hết mọi rủi ro.\n\nQuản trị dự án xuất sắc là người:\n1. Biết khóa chặt những biến số rủi ro cốt lõi nhất.\n2. Rèn luyện cho hệ thống một SỨC ĐỀ KHÁNG để bật dậy nhanh nhất sau mỗi lần vấp ngã.",
        "image": "DS_Cuoiky/call_to_action_1783750036484.jpg",
        "notes": "Xin kết lại bài thuyết trình của nhóm bằng một thông điệp từ chính dữ liệu..."
    },
    {
        "type": "title",
        "title": "Q&A",
        "subtitle": "Xin trân trọng cảm ơn Thầy và các bạn đã lắng nghe!\n\n(Vui lòng xem file Jupyter Notebook và Dataset đính kèm)",
        "notes": "Cảm ơn Thầy và các bạn đã lắng nghe câu chuyện dữ liệu của nhóm chúng em..."
    }
]

prs = Presentation()

for slide_data in slides_data:
    if slide_data["type"] == "title":
        slide_layout = prs.slide_layouts[0] # Title
        slide = prs.slides.add_slide(slide_layout)
        title = slide.shapes.title
        subtitle = slide.placeholders[1]
        
        title.text = slide_data["title"]
        subtitle.text = slide_data["subtitle"]
        
        if "image" in slide_data and os.path.exists(slide_data["image"]):
            # Add image to title slide (centered above title if possible, or right side)
            left = Inches(3.5)
            top = Inches(0.2)
            height = Inches(2.5)
            slide.shapes.add_picture(slide_data["image"], left, top, height=height)
            
    elif slide_data["type"] == "content":
        slide_layout = prs.slide_layouts[1] # Content
        slide = prs.slides.add_slide(slide_layout)
        title = slide.shapes.title
        body = slide.placeholders[1]
        
        title.text = slide_data["title"]
        body.text = slide_data["body"]
        
    elif slide_data["type"] == "content_with_image":
        slide_layout = prs.slide_layouts[1] # Title and Content
        slide = prs.slides.add_slide(slide_layout)
        title = slide.shapes.title
        body = slide.placeholders[1]
        
        title.text = slide_data["title"]
        body.text = slide_data["body"]
        
        # Resize body text box to make room for image
        body.width = Inches(3.5)
        
        if "image" in slide_data and os.path.exists(slide_data["image"]):
            # Add image
            top = Inches(2.0)
            left = Inches(4.0)
            height = Inches(4.5)
            slide.shapes.add_picture(slide_data["image"], left, top, height=height)

    # Add speaker notes
    notes_slide = slide.notes_slide
    text_frame = notes_slide.notes_text_frame
    if "notes" in slide_data:
        text_frame.text = slide_data["notes"]

prs.save(pptx_path)
print(f"Đã cập nhật file PowerPoint với FULL HÌNH ẢNH (Biểu đồ + Minh họa): {pptx_path}")
