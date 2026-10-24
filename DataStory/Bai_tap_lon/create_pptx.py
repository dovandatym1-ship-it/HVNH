import sys
import subprocess

# Ensure python-pptx is installed
try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.text import PP_ALIGN
    from pptx.dml.color import RGBColor
except ImportError:
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-pptx"])
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.text import PP_ALIGN
    from pptx.dml.color import RGBColor

# Kịch bản các slide
slides_data = [
    {
        "type": "title",
        "title": "Không sợ 'Lỗ hổng' phát sinh, chỉ sợ thiếu 'Khả năng phục hồi'",
        "subtitle": "Giải pháp tối ưu tiến độ dự án xây dựng\n(Phân tích từ 6.000 dự án - Công ty CP Tập đoàn Công nghiệp Việt)",
        "notes": "Chào thầy và các bạn. Quản trị dự án không phải là chạy theo ảo vọng 0% rủi ro. Hôm nay, nhóm chúng em sẽ chứng minh: Không sợ dự án phát sinh lỗ hổng, chỉ sợ bộ máy quản lý thiếu vắng khả năng phục hồi."
    },
    {
        "type": "content",
        "title": "1. Bối cảnh (The Context): Bức tranh 40% trễ hạn",
        "body": "• Giai đoạn 2022-2025: Triển khai 6.000 dự án lớn nhỏ.\n• Thực trạng:\n  - 60% dự án hoàn thành xuất sắc.\n  - 40% dự án BỊ TRỄ HẠN.\n• Vấn đề: Khi trễ hạn, dự án trượt tiến độ trung bình lên tới gần 2 tháng. Đâu là lỗ hổng khổng lồ nuốt chửng các dự án này?",
        "notes": "Hãy cùng bước vào câu chuyện của Tập đoàn Công nghiệp Việt..."
    },
    {
        "type": "content",
        "title": "2. Nút thắt (The Conflict): Xóa bỏ định kiến",
        "body": "• Các định kiến truyền thống (Những kẻ chịu trận oan uổng):\n  - Thời tiết xấu\n  - Giải phóng mặt bằng\n  - Nhân công nghỉ việc\n• SỰ THẬT TỪ DỮ LIỆU:\n  - Các nguyên nhân trên chỉ đóng góp chưa tới 15%.\n  - Kẻ kéo lùi dự án tàn bạo nhất lộ diện: THAY ĐỔI THIẾT KẾ!",
        "notes": "Khi ban giám đốc đổ lỗi cho ông trời, dữ liệu Bar Chart đã bác bỏ hoàn toàn..."
    },
    {
        "type": "content",
        "title": "3. Điểm bùng nổ (The Aha! Moment)",
        "body": "Lỗ hổng thực sự mang tên: THAY ĐỔI THIẾT KẾ\n\n• Tương quan tỷ lệ thuận tuyệt đối (Scatter Plot):\n  - Càng sửa bản vẽ, dự án càng trượt dài!\n  - Cứ mỗi một lần chủ đầu tư 'đổi ý' trên giấy, quỹ đạo dự án trượt dốc.\n• Nó không phải lỗ hổng ngoại cảnh, mà là lỗ hổng do con người.",
        "notes": "Và đây chính là khoảnh khắc 'Aha!' – điểm sáng nhất trong toàn bộ bài phân tích..."
    },
    {
        "type": "content",
        "title": "4. Phân tích sâu (Deep Dive): Hiệu ứng Hòn tuyết lăn",
        "body": "Cơ chế tàn phá của việc thay đổi thiết kế:\n\n• Thay đổi bản vẽ -> Chuỗi cung ứng vật tư đứt gãy -> Công nhân dừng việc chờ đợi.\n• SỰ THIẾU HỤT KHẢ NĂNG PHỤC HỒI TIẾN ĐỘ:\n  - Thay vì dập tắt rủi ro linh hoạt, bộ máy quản lý để mặc thả trôi.\n  - Một lỗ hổng thiết kế nhỏ ban đầu tích tụ thành quả cầu tuyết khổng lồ.",
        "notes": "Tại sao lấy cục tẩy xóa 1 bức tường lại tàn phá khủng khiếp đến thế?..."
    },
    {
        "type": "content",
        "title": "5. Giải pháp Kép: Vá Lỗ hổng (Phần 1)",
        "body": "VÁ LỖ HỔNG TỪ PHÒNG THIẾT KẾ\n\n• Kỷ luật 'Design Freeze' (Khóa bản vẽ):\n  - Thiết lập một mốc thời gian tử thần trong hợp đồng.\n  - Mọi thay đổi sau mốc này bị cấm hoặc phải bị kiểm soát chặt chẽ.\n• Đừng chạy đua trên bùn lầy, hãy chốt chặn từ phòng điều hòa!",
        "notes": "Để giải cứu Tập đoàn, chúng em đề xuất 2 bước hành động. Bước 1: Vá lỗ hổng..."
    },
    {
        "type": "content",
        "title": "6. Giải pháp Kép: Tăng Phục hồi (Phần 2)",
        "body": "TĂNG CƯỜNG KHẢ NĂNG PHỤC HỒI (SCHEDULE RESILIENCE)\n\n• Quy trình 'Change Order' (Biên lai thay đổi):\n  - Yêu cầu đổi thiết kế = Tiền phạt + Thời gian gia hạn.\n• Kích hoạt nén tiến độ:\n  - Bơm ngân sách làm song song (Fast-tracking).\n  - Dập tắt độ trễ trong vòng 48h để chặn đứng quả cầu tuyết.",
        "notes": "Nhưng nếu khách hàng vẫn kiên quyết bắt ép đổi thiết kế thì sao? Đó là lúc ta cần Bước 2..."
    },
    {
        "type": "content",
        "title": "7. Hỗ trợ Công nghệ: Áp dụng mô hình BIM 3D",
        "body": "Tương lai của quản lý thiết kế:\n\n• Chuyển dịch từ bản vẽ 2D sang mô hình 3D (BIM).\n• Trực quan hóa công trình trước khi khởi công.\n• Giảm thiểu tối đa sự dao động tâm lý của khách hàng, hạn chế thay đổi thiết kế về sau.",
        "notes": "Ngoài pháp lý hợp đồng, việc chuyển dịch sang mô hình 3D (BIM) cũng là một vũ khí..."
    },
    {
        "type": "content",
        "title": "8. Lời kêu gọi (Call to Action)",
        "body": "Quản trị dự án xuất sắc không phải là người giăng lưới bắt hết mọi rủi ro.\n\nQuản trị dự án xuất sắc là người:\n1. Biết khóa chặt những biến số rủi ro cốt lõi nhất.\n2. Rèn luyện cho hệ thống một SỨC ĐỀ KHÁNG để bật dậy nhanh nhất sau mỗi lần vấp ngã.",
        "notes": "Xin kết lại bài thuyết trình của nhóm bằng một thông điệp từ chính dữ liệu..."
    },
    {
        "type": "title",
        "title": "Q&A",
        "subtitle": "Xin trân trọng cảm ơn Thầy và các bạn đã lắng nghe!\n\n(Vui lòng xem file Jupyter Notebook và Dataset đính kèm)",
        "notes": "Cảm ơn Thầy và các bạn đã lắng nghe câu chuyện dữ liệu của nhóm chúng em..."
    }
]

def create_presentation():
    prs = Presentation()
    
    for slide_data in slides_data:
        if slide_data["type"] == "title":
            slide_layout = prs.slide_layouts[0] # Title slide
            slide = prs.slides.add_slide(slide_layout)
            title = slide.shapes.title
            subtitle = slide.placeholders[1]
            
            title.text = slide_data["title"]
            subtitle.text = slide_data["subtitle"]
            
        elif slide_data["type"] == "content":
            slide_layout = prs.slide_layouts[1] # Title and Content slide
            slide = prs.slides.add_slide(slide_layout)
            title = slide.shapes.title
            body = slide.placeholders[1]
            
            title.text = slide_data["title"]
            body.text = slide_data["body"]
        
        # Add speaker notes
        notes_slide = slide.notes_slide
        text_frame = notes_slide.notes_text_frame
        text_frame.text = slide_data["notes"]

    # Save presentation
    prs.save("DS_Cuoiky/Slide_Thuyet_Trinh_Chi_Tiet.pptx")
    print("Đã tạo thành công file PowerPoint: DS_Cuoiky/Slide_Thuyet_Trinh_Chi_Tiet.pptx")

if __name__ == "__main__":
    create_presentation()
