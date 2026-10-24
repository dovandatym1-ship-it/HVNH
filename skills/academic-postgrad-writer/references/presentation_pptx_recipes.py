"""
Antigravity Academic Writer - Python-pptx Modular Recipe Library for Executive Academic Presentations
Implements 16:9 Widescreen, 4-Layer Architecture, Action Titles, Card Grid & Speaker Notes
"""
import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

# ==============================================================================
# HỆ MÀU NHẬN DIỆN THƯƠNG HIỆU HỌC VIỆN NGÂN HÀNG & EXECUTIVE DASHBOARD
# ==============================================================================
C_HVNH_NAVY    = RGBColor(10, 37, 64)       # #0A2540: Xanh Navy đậm BAV
C_HVNH_ORANGE  = RGBColor(234, 88, 12)      # #EA580C: Cam năng lượng BAV
C_HVNH_GOLD    = RGBColor(245, 158, 11)     # #F59E0B: Vàng điểm nhấn
C_BLUE_CYAN    = RGBColor(2, 132, 199)      # #0284C7: Xanh Cyan chỉ số phụ
C_ACCENT_RED   = RGBColor(220, 38, 38)      # #DC2626: Đỏ cảnh báo rủi ro
C_ACCENT_GREEN = RGBColor(16, 185, 129)     # #10B981: Xanh lá thành công

C_CANVAS_BG    = RGBColor(248, 250, 252)    # #F8FAFC: Nền canvas Slate 50
C_CARD_BG      = RGBColor(255, 255, 255)    # #FFFFFF: Nền thẻ Card trắng
C_CARD_BORDER  = RGBColor(226, 232, 240)    # #E2E8F0: Viền thẻ Slate 200
C_TEXT_DARK    = RGBColor(15, 23, 42)       # #0F172A: Màu chữ Slate 900
C_TEXT_MUTED   = RGBColor(100, 116, 139)    # #64748B: Màu chữ chú thích Slate 500
C_WHITE        = RGBColor(255, 255, 255)

FONT_PRIMARY = 'Calibri'
FONT_TITLE   = 'Calibri'

def create_executive_deck():
    """Khởi tạo Deck chuẩn Widescreen 16:9 (13.333 x 7.5 inches)"""
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs

def add_blank_slide(prs, bg_color=C_CANVAS_BG):
    """Tạo slide trống với màu nền canvas tinh chỉnh"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = bg_color
    bg.line.fill.background()
    return slide

def add_action_header(slide, tag_text, title_text, page_num=None, total_pages=20):
    """
    Tạo Header chuẩn Action Title của EasySlides / McKinsey:
    - Pill Badge ở trên (Tag)
    - Action Title đậm nét đúc kết thông điệp chính ở dưới
    - Page Number đếm trang ở góc phải
    """
    # Pill Badge Tag
    tag_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.42), Inches(3.2), Inches(0.32))
    tag_box.fill.solid()
    tag_box.fill.fore_color.rgb = C_HVNH_NAVY
    tag_box.line.fill.background()
    tf_tag = tag_box.text_frame
    tf_tag.word_wrap = True
    tf_tag.margin_left = Inches(0.1)
    tf_tag.margin_top = Inches(0.04)
    p_tag = tf_tag.paragraphs[0]
    p_tag.alignment = PP_ALIGN.CENTER
    r_tag = p_tag.add_run()
    r_tag.text = tag_text.upper()
    r_tag.font.name = FONT_PRIMARY
    r_tag.font.size = Pt(9.5)
    r_tag.font.bold = True
    r_tag.font.color.rgb = C_WHITE

    # Action Title Text
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.78), Inches(10.5), Inches(0.75))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0)
    tf.margin_top = Inches(0)
    p_title = tf.paragraphs[0]
    p_title.alignment = PP_ALIGN.LEFT
    r_title = p_title.add_run()
    r_title.text = title_text
    r_title.font.name = FONT_TITLE
    r_title.font.size = Pt(20)
    r_title.font.bold = True
    r_title.font.color.rgb = C_HVNH_NAVY

    # Page Number
    if page_num:
        tb_p = slide.shapes.add_textbox(Inches(11.5), Inches(0.42), Inches(1.2), Inches(0.35))
        p_p = tb_p.text_frame.paragraphs[0]
        p_p.alignment = PP_ALIGN.RIGHT
        r_p = p_p.add_run()
        r_p.text = f"{page_num:02d} / {total_pages:02d}"
        r_p.font.name = FONT_PRIMARY
        r_p.font.size = Pt(11)
        r_p.font.bold = True
        r_p.font.color.rgb = C_TEXT_MUTED

def add_card_box(slide, left, top, width, height, bg_color=C_CARD_BG, border_color=C_CARD_BORDER, border_width=Pt(1.2)):
    """Tạo khung Card bo góc viền tinh tế"""
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    if border_color:
        card.line.color.rgb = border_color
        card.line.width = border_width
    else:
        card.line.fill.background()
    return card

def add_metric_card(slide, left, top, width, height, lead_kpi, kpi_label, desc_lines, highlight_color=C_HVNH_ORANGE):
    """Tạo thẻ Metric với số liệu lớn (Lead KPI) nổi bật"""
    add_card_box(slide, left, top, width, height)
    
    # Top Accent Strip
    strip = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left + Inches(0.15), top + Inches(0.15), width - Inches(0.3), Inches(0.06))
    strip.fill.solid()
    strip.fill.fore_color.rgb = highlight_color
    strip.line.fill.background()

    tb = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.3), width - Inches(0.4), height - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    
    # KPI Number
    p_kpi = tf.paragraphs[0]
    p_kpi.alignment = PP_ALIGN.LEFT
    r_kpi = p_kpi.add_run()
    r_kpi.text = lead_kpi
    r_kpi.font.name = FONT_TITLE
    r_kpi.font.size = Pt(30)
    r_kpi.font.bold = True
    r_kpi.font.color.rgb = highlight_color

    # KPI Label
    p_lbl = tf.add_paragraph()
    p_lbl.space_before = Pt(2)
    p_lbl.space_after = Pt(6)
    r_lbl = p_lbl.add_run()
    r_lbl.text = kpi_label
    r_lbl.font.name = FONT_PRIMARY
    r_lbl.font.size = Pt(12)
    r_lbl.font.bold = True
    r_lbl.font.color.rgb = C_HVNH_NAVY

    # Descriptions
    for line in desc_lines:
        p_d = tf.add_paragraph()
        p_d.space_before = Pt(2)
        r_d = p_d.add_run()
        r_d.text = f"• {line}"
        r_d.font.name = FONT_PRIMARY
        r_d.font.size = Pt(10.5)
        r_d.font.color.rgb = C_TEXT_DARK

def add_chart_deep_dive_60_40(slide, img_path, left_notes_title, note_items):
    """Bố cục Chart Deep-Dive 60/40: 60% biểu đồ hình ảnh bên trái, 40% phân tích bên phải"""
    # 60% Image on Left
    if os.path.exists(img_path):
        add_card_box(slide, Inches(0.8), Inches(1.65), Inches(7.4), Inches(5.2))
        slide.shapes.add_picture(img_path, Inches(0.95), Inches(1.8), width=Inches(7.1))

    # 40% Card on Right
    add_card_box(slide, Inches(8.4), Inches(1.65), Inches(4.1), Inches(5.2))
    tb = slide.shapes.add_textbox(Inches(8.6), Inches(1.85), Inches(3.7), Inches(4.8))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p_hdr = tf.paragraphs[0]
    r_hdr = p_hdr.add_run()
    r_hdr.text = left_notes_title
    r_hdr.font.name = FONT_TITLE
    r_hdr.font.size = Pt(14)
    r_hdr.font.bold = True
    r_hdr.font.color.rgb = C_HVNH_NAVY

    for tag, desc in note_items:
        p_item = tf.add_paragraph()
        p_item.space_before = Pt(8)
        
        r_tag = p_item.add_run(f"📌 {tag}: ")
        r_tag.font.name = FONT_PRIMARY
        r_tag.font.size = Pt(11)
        r_tag.font.bold = True
        r_tag.font.color.rgb = C_HVNH_ORANGE
        
        r_desc = p_item.add_run(desc)
        r_desc.font.name = FONT_PRIMARY
        r_desc.font.size = Pt(10.5)
        r_desc.font.color.rgb = C_TEXT_DARK

def add_process_pipeline_4(slide, steps_data):
    """
    Tạo quy trình 4 bước ngang (Horizontal 4-Step Pipeline)
    steps_data = [("01", "TIỀN XỬ LÝ", ["Loại bỏ 2 cột rò rỉ", "Điền khuyết thiếu"]), ...]
    """
    card_w = Inches(2.75)
    card_h = Inches(5.1)
    top_pos = Inches(1.7)
    
    for i, (num_str, title, lines) in enumerate(steps_data):
        left_pos = Inches(0.8) + i * Inches(2.95)
        add_card_box(slide, left_pos, top_pos, card_w, card_h)
        
        # Step Badge
        badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_pos + Inches(0.18), top_pos + Inches(0.18), Inches(0.7), Inches(0.4))
        badge.fill.solid()
        badge.fill.fore_color.rgb = C_HVNH_ORANGE if i == 0 else C_HVNH_NAVY
        badge.line.fill.background()
        tf_b = badge.text_frame
        p_b = tf_b.paragraphs[0]
        p_b.alignment = PP_ALIGN.CENTER
        r_b = p_b.add_run()
        r_b.text = num_str
        r_b.font.bold = True
        r_b.font.color.rgb = C_WHITE
        
        # Text Content
        tb = slide.shapes.add_textbox(left_pos + Inches(0.18), top_pos + Inches(0.7), card_w - Inches(0.36), card_h - Inches(0.9))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p_t = tf.paragraphs[0]
        r_t = p_t.add_run()
        r_t.text = title
        r_t.font.bold = True
        r_t.font.size = Pt(13)
        r_t.font.color.rgb = C_HVNH_NAVY
        
        for line in lines:
            p_l = tf.add_paragraph()
            p_l.space_before = Pt(6)
            r_l = p_l.add_run()
            r_l.text = f"• {line}"
            r_l.font.size = Pt(10.5)
            r_l.font.color.rgb = C_TEXT_DARK

def add_persona_profile_cards(slide, personas_data):
    """
    Tạo 3 thẻ Chân dung khách hàng Personas
    personas_data = [("1. Silent Defector", "55% Rời bỏ", "Dư nợ $0, quẹt thẻ <45 lần", "E-voucher cuối tuần, hoàn tiền 5%"), ...]
    """
    card_w = Inches(3.7)
    card_h = Inches(5.1)
    top_pos = Inches(1.7)
    
    for i, (name, share, features, action) in enumerate(personas_data):
        left_pos = Inches(0.8) + i * Inches(3.95)
        add_card_box(slide, left_pos, top_pos, card_w, card_h)
        
        # Header Box in Card
        hdr_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_pos + Inches(0.15), top_pos + Inches(0.15), card_w - Inches(0.3), Inches(0.85))
        hdr_box.fill.solid()
        hdr_box.fill.fore_color.rgb = C_HVNH_NAVY if i != 0 else C_HVNH_ORANGE
        hdr_box.line.fill.background()
        tf_h = hdr_box.text_frame
        p_h0 = tf_h.paragraphs[0]
        p_h0.alignment = PP_ALIGN.CENTER
        r_h0 = p_h0.add_run()
        r_h0.text = name
        r_h0.font.bold = True
        r_h0.font.size = Pt(12)
        r_h0.font.color.rgb = C_WHITE
        
        p_h1 = tf_h.add_paragraph()
        p_h1.alignment = PP_ALIGN.CENTER
        r_h1 = p_h1.add_run()
        r_h1.text = share
        r_h1.font.bold = True
        r_h1.font.size = Pt(11)
        r_h1.font.color.rgb = C_WHITE
        
        # Details
        tb = slide.shapes.add_textbox(left_pos + Inches(0.2), top_pos + Inches(1.15), card_w - Inches(0.4), card_h - Inches(1.3))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p_f_lbl = tf.paragraphs[0]
        r_f_lbl = p_f_lbl.add_run("🔍 Đặc trưng hành vi:")
        r_f_lbl.font.bold = True
        r_f_lbl.font.size = Pt(11)
        r_f_lbl.font.color.rgb = C_HVNH_NAVY
        
        p_f_val = tf.add_paragraph()
        p_f_val.space_before = Pt(2)
        p_f_val.space_after = Pt(10)
        r_f_val = p_f_val.add_run(features)
        r_f_val.font.size = Pt(10.5)
        r_f_val.font.color.rgb = C_TEXT_DARK
        
        p_a_lbl = tf.add_paragraph()
        r_a_lbl = p_a_lbl.add_run("🎯 Hành động can thiệp:")
        r_a_lbl.font.bold = True
        r_a_lbl.font.size = Pt(11)
        r_a_lbl.font.color.rgb = C_HVNH_ORANGE
        
        p_a_val = tf.add_paragraph()
        p_a_val.space_before = Pt(2)
        r_a_val = p_a_val.add_run(action)
        r_a_val.font.size = Pt(10.5)
        r_a_val.font.color.rgb = C_TEXT_DARK

def add_speaker_notes(slide, notes_text):
    """Chèn lời thoại thuyết trình (Speaker Notes) tích hợp sẵn trong Slide"""
    notes_slide = slide.notes_slide
    text_frame = notes_slide.notes_text_frame
    text_frame.text = notes_text
