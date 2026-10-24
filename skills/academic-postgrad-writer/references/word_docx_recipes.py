"""
Academic Postgraduate Writer — Complete Python-docx Recipes for Postgraduate Academic Reports
v2 — Profile-based standards engine:
  * hvnh-thesis    : Luận văn Thạc sĩ HVNH (Quy định về trình bày luận văn — Khoa Sau đại học HVNH)
                     Lề 3.5/3.0/3.5/2.0 cm • Giãn dòng 1.5 • Số trang giữa-phía-trên đếm từ MỞ ĐẦU
                     Trích dẫn số [15, tr.314] • Caption gắn số chương (Hình 3.4)
  * hvnh-khoaluan  : Khóa luận / Chuyên đề tốt nghiệp (giãn dòng 1.3–1.5, mặc định 1.5)
  * nd30-report    : Báo cáo hành chính theo ND 30/2020/NĐ-CP (legacy profile của v1)
Xem chi tiết chuẩn gốc: references/hvnh_thesis_standard.md
"""
import os
import docx
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_SECTION
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

# ==============================================================================
# HỆ MÀU CHUẨN HỌC VIỆN NGÂN HÀNG & BÁO CÁO TÀI CHÍNH KHOA HỌC
# ==============================================================================
COLOR_NAVY_HVNH   = RGBColor(10, 37, 64)       # #0A2540: Xanh Navy đậm HVNH
COLOR_BLUE_SEC    = RGBColor(31, 77, 120)      # #1F4D78: Xanh Slate tiêu đề cấp 2
COLOR_ACCENT_BLUE = RGBColor(46, 116, 181)     # #2E74B5: Xanh tiêu đề cấp 3
COLOR_ORANGE_HVNH = RGBColor(234, 88, 12)      # #EA580C: Cam năng lượng BAV
COLOR_DARK_TEXT   = RGBColor(15, 23, 42)       # #0F172A: Màu chữ Slate 900
COLOR_MUTED_TEXT  = RGBColor(100, 116, 139)    # #64748B: Màu chữ chú thích Slate 500

HEX_NAVY_BG    = "0A2540"
HEX_BLUE_HDR   = "17365D"
HEX_ZEBRA_BG   = "F8FAFC"
HEX_CALLOUT_BG = "F1F5F9"
HEX_BORDER     = "CBD5E1"
HEX_ORANGE     = "EA580C"

# ==============================================================================
# PROFILE ĐỊNH DẠNG — chọn đúng chuẩn trước khi soạn thảo
# ==============================================================================
PROFILES = {
    "hvnh-thesis": {
        "name": "Luận văn Thạc sĩ HVNH (Quy định Khoa Sau đại học)",
        "margins_cm": {"top": 3.5, "bottom": 3.0, "left": 3.5, "right": 2.0},
        "line_spacing": 1.5,
        "page_number_pos": "top-center",
        "citation_style": "numeric-bracketed",
        "caption_numbering": "chapter-based",
    },
    "hvnh-khoaluan": {
        "name": "Khóa luận / Chuyên đề tốt nghiệp HVNH",
        "margins_cm": {"top": 3.5, "bottom": 3.0, "left": 3.5, "right": 2.0},
        "line_spacing": 1.5,
        "page_number_pos": "top-center",
        "citation_style": "numeric-bracketed",
        "caption_numbering": "chapter-based",
    },
    "nd30-report": {
        "name": "Báo cáo hành chính ND 30/2020/NĐ-CP (legacy v1)",
        "margins_cm": {"top": 2.5, "bottom": 2.0, "left": 3.0, "right": 2.0},
        "line_spacing": 1.35,
        "page_number_pos": "footer",
        "citation_style": "author-date",
        "caption_numbering": "flat",
    },
}
CURRENT_PROFILE = PROFILES["hvnh-thesis"]

def get_profile(name):
    """Lấy cấu hình profile theo tên; ném lỗi rõ ràng nếu sai tên."""
    if name not in PROFILES:
        raise ValueError(
            f"Profile '{name}' không tồn tại. Các profile khả dụng: {list(PROFILES)}"
        )
    return PROFILES[name]

# ==============================================================================
# CẤU HÌNH TÀI LIỆU THEO PROFILE
# ==============================================================================
def configure_academic_document(doc, profile="hvnh-thesis"):
    """
    Cấu hình khổ giấy A4 + lề + giãn dòng theo profile đã chọn và khởi tạo hệ thống
    Built-in Heading Styles cho Word Navigation Pane.
    Mặc định 'hvnh-thesis': Trái 3.5cm, Phải 2.0cm, Trên 3.5cm, Dưới 3.0cm, giãn dòng 1.5.
    """
    global CURRENT_PROFILE
    CURRENT_PROFILE = get_profile(profile)

    sec = doc.sections[0]
    m = CURRENT_PROFILE["margins_cm"]
    sec.top_margin = Cm(m["top"])
    sec.bottom_margin = Cm(m["bottom"])
    sec.left_margin = Cm(m["left"])
    sec.right_margin = Cm(m["right"])
    sec.page_width = Cm(21.0)      # 210 mm (A4)
    sec.page_height = Cm(29.7)     # 297 mm (A4)

    ls = CURRENT_PROFILE["line_spacing"]

    # Style Normal (Body Text)
    s_norm = doc.styles['Normal']
    s_norm.font.name = 'Times New Roman'
    s_norm.font.size = Pt(13)
    s_norm.font.color.rgb = COLOR_DARK_TEXT
    s_norm.paragraph_format.line_spacing = ls
    s_norm.paragraph_format.space_after = Pt(6)
    s_norm.paragraph_format.space_before = Pt(0)
    s_norm.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # Built-in Heading 1 (Word Navigation Pane Level 0)
    s_h1 = doc.styles['Heading 1']
    s_h1.font.name = 'Times New Roman'
    s_h1.font.size = Pt(15.0)
    s_h1.font.bold = True
    s_h1.font.color.rgb = COLOR_NAVY_HVNH
    s_h1.paragraph_format.space_before = Pt(16)
    s_h1.paragraph_format.space_after = Pt(8)
    s_h1.paragraph_format.keep_with_next = True
    s_h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    s_h1.paragraph_format.first_line_indent = Cm(0)

    # Built-in Heading 2 (Word Navigation Pane Level 1)
    s_h2 = doc.styles['Heading 2']
    s_h2.font.name = 'Times New Roman'
    s_h2.font.size = Pt(13.5)
    s_h2.font.bold = True
    s_h2.font.color.rgb = COLOR_BLUE_SEC
    s_h2.paragraph_format.space_before = Pt(14)
    s_h2.paragraph_format.space_after = Pt(6)
    s_h2.paragraph_format.keep_with_next = True
    s_h2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    s_h2.paragraph_format.first_line_indent = Cm(0)

    # Built-in Heading 3 (Word Navigation Pane Level 2)
    s_h3 = doc.styles['Heading 3']
    s_h3.font.name = 'Times New Roman'
    s_h3.font.size = Pt(13.0)
    s_h3.font.bold = True
    s_h3.font.color.rgb = COLOR_ACCENT_BLUE
    s_h3.paragraph_format.space_before = Pt(10)
    s_h3.paragraph_format.space_after = Pt(4)
    s_h3.paragraph_format.keep_with_next = True
    s_h3.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    s_h3.paragraph_format.first_line_indent = Cm(0)

# ==============================================================================
# SỐ TRANG GIỮA-PHÍA-TRÊN, ĐẾM TỪ MỞ ĐẦU (chuẩn luận văn HVNH)
# ==============================================================================
def _set_page_number_start(section, start=1):
    """Ghi <w:pgNumType w:start='N'/> vào sectPr theo đúng vị trí schema (trước <w:cols>)."""
    sectPr = section._sectPr
    for el in sectPr.findall(qn('w:pgNumType')):
        sectPr.remove(el)
    pg = parse_xml('<w:pgNumType %s w:start="%d"/>' % (nsdecls('w'), start))
    cols = sectPr.find(qn('w:cols'))
    if cols is not None:
        cols.addprevious(pg)
    else:
        sectPr.append(pg)

def apply_page_numbers_top_center(doc, start_value=1):
    """
    Gọi NGAY TRƯỚC khi thêm tiêu đề 'MỞ ĐẦU'.
    Tạo section mới (ngắt trang): mọi section trước đó (bìa, cam đoan, mục lục...)
    không hiển thị số trang; từ section này trở đi, số trang nằm ở GIỮA PHÍA TRÊN
    và bắt đầu lại từ start_value (=1 theo quy định).
    Trả về section thân bài để tiện kiểm chứng.
    """
    body_sec = doc.add_section(WD_SECTION.NEW_PAGE)
    body_sec.header.is_linked_to_previous = False
    hdr_p = body_sec.header.paragraphs[0]
    hdr_p.paragraph_format.first_line_indent = Cm(0)
    hdr_p.paragraph_format.space_after = Pt(6)
    hdr_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_xml_page_field(hdr_p, "PAGE", default_val=str(start_value))
    _set_page_number_start(body_sec, start=start_value)
    return body_sec

# ==============================================================================
# TIỆN ÍCH XML CƠ SỞ
# ==============================================================================
def set_cell_background(cell, hex_color):
    """Thiết lập màu nền ô bảng (Shading)"""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=90, bottom=90, left=110, right=110):
    """Thiết lập khoảng đệm trong ô bảng (dxa: 20 dxa = 1 pt)"""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_xml_page_field(p, field_name="PAGE", default_val="1"):
    """Thêm mã trường đếm trang động OpenXML (PAGE / NUMPAGES) kèm fallback cached value"""
    run = p.add_run()
    fldChar1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instrText = parse_xml(r'<w:instrText %s xml:space="preserve"> %s </w:instrText>' % (nsdecls('w'), field_name))
    fldChar2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    cached = parse_xml(r'<w:t %s>%s</w:t>' % (nsdecls('w'), default_val))
    fldChar3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(cached)
    run._r.append(fldChar3)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(10)
    run.font.color.rgb = COLOR_MUTED_TEXT

# ==============================================================================
# HEADING & NỘI DUNG THÂN BÀI
# ==============================================================================
def add_chapter_heading(doc, text, is_first_chapter=False):
    """Thêm Heading 1 kèm ngắt trang trước mỗi chương (Page break per chapter).
    Lưu ý chuẩn HVNH: tên chương IN HOA, không kết thúc bằng dấu chấm."""
    if not is_first_chapter:
        doc.add_page_break()
    p = doc.add_paragraph(text, style='Heading 1')
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Cm(0)
    return p

def add_section_heading(doc, text):
    """Thêm Heading 2 (Mục cấp 1)"""
    p = doc.add_paragraph(text, style='Heading 2')
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Cm(0)
    return p

def add_subsection_heading(doc, text):
    """Thêm Heading 3 (Mục cấp 2). Chuẩn HVNH: tối đa 4 cấp số (3.1.2.1);
    cấm tiểu mục mồ côi — phải có ít nhất 2 tiểu mục trong cùng nhóm."""
    p = doc.add_paragraph(text, style='Heading 3')
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Cm(0)
    return p

def add_body_paragraph(doc, text, indent=True, keep_with_next=False):
    """Thêm đoạn văn nội dung căn đều (Justified) thụt đầu dòng 1.27cm"""
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = CURRENT_PROFILE["line_spacing"]
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.keep_with_next = keep_with_next
    if indent:
        p.paragraph_format.first_line_indent = Inches(0.5)  # 1.27 cm chuẩn
    else:
        p.paragraph_format.first_line_indent = Cm(0)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(13)
    run.font.color.rgb = COLOR_DARK_TEXT
    return p

def add_block_quote(doc, text, italic=False):
    """Trích dẫn dài hơn 2 câu hoặc 4 dòng (chuẩn HVNH §2.5):
    tách thành đoạn riêng, lề trái lùi vào thêm 2cm, KHÔNG dùng ngoặc kép."""
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = CURRENT_PROFILE["line_spacing"]
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Cm(2)
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(13)
    run.font.italic = italic
    run.font.color.rgb = COLOR_DARK_TEXT
    return p

def add_chapter_conclusion(doc, chapter_no, summary_text):
    """Tóm tắt chương (kết luận chương) — bắt buộc ở cuối mỗi chương theo quy định.
    Không nhất thiết sang trang riêng."""
    p_h = doc.add_paragraph(f"TÓM TẮT CHƯƠNG {chapter_no}", style='Heading 2')
    p_h.paragraph_format.space_before = Pt(14)
    p_h.paragraph_format.keep_with_next = True
    return add_body_paragraph(doc, summary_text)

def add_numeric_citation(paragraph, refs, pages=None, before_text=None):
    """Trích dẫn số chuẩn HVNH trong ngoặc vuông, gắn vào paragraph hiện có.
    refs=[15]              -> [15]
    refs=[15], pages='314' -> [15, tr.314]
    refs=[19,25]           -> [19], [25]  (tăng dần, mỗi nguồn một ngoặc)
    """
    refs = sorted(refs)
    if len(refs) == 1:
        marker = f"[{refs[0]}, tr.{pages}]" if pages else f"[{refs[0]}]"
    else:
        marker = ", ".join(f"[{r}]" for r in refs)
    if before_text:
        r_pre = paragraph.add_run(before_text)
        r_pre.font.name = 'Times New Roman'
        r_pre.font.size = Pt(13)
        r_pre.font.color.rgb = COLOR_DARK_TEXT
    r = paragraph.add_run(marker)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(13)
    r.font.color.rgb = COLOR_DARK_TEXT
    return paragraph

def add_numbered_item(doc, number_str, lead_text, body_text):
    """Thêm mục đánh số với hanging indent 0.3 inch"""
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = min(CURRENT_PROFILE["line_spacing"], 1.32)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.first_line_indent = Inches(-0.3)

    r_num = p.add_run(number_str + " ")
    r_num.bold = True
    r_num.font.name = 'Times New Roman'
    r_num.font.size = Pt(13)
    r_num.font.color.rgb = COLOR_NAVY_HVNH

    r_lead = p.add_run(lead_text + " ")
    r_lead.bold = True
    r_lead.font.name = 'Times New Roman'
    r_lead.font.size = Pt(13)
    r_lead.font.color.rgb = COLOR_NAVY_HVNH

    r_body = p.add_run(body_text)
    r_body.font.name = 'Times New Roman'
    r_body.font.size = Pt(13)
    r_body.font.color.rgb = COLOR_DARK_TEXT
    return p

def add_bullet_item(doc, lead_text, body_text):
    """Thêm gạch đầu dòng căn Trái (Left) chống dãn chữ"""
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = min(CURRENT_PROFILE["line_spacing"], 1.3)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.first_line_indent = Inches(-0.2)

    r_bullet = p.add_run("• ")
    r_bullet.bold = True
    r_bullet.font.name = 'Times New Roman'
    r_bullet.font.size = Pt(13)
    r_bullet.font.color.rgb = COLOR_NAVY_HVNH

    if lead_text:
        r_lead = p.add_run(lead_text + " ")
        r_lead.bold = True
        r_lead.font.name = 'Times New Roman'
        r_lead.font.size = Pt(13)
        r_lead.font.color.rgb = COLOR_NAVY_HVNH

    r_body = p.add_run(body_text)
    r_body.font.name = 'Times New Roman'
    r_body.font.size = Pt(13)
    r_body.font.color.rgb = COLOR_DARK_TEXT
    return p

def add_callout_box(doc, title, text, is_warning=False):
    """Tạo hộp Callout với viền trái màu nổi bật, nền xám nhạt và căn Trái chống dãn chữ"""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_background(cell, HEX_CALLOUT_BG)
    set_cell_margins(cell, top=120, bottom=120, left=160, right=140)

    border_color = HEX_ORANGE if is_warning else HEX_BLUE_HDR
    icon = "🚨" if is_warning else "💡"

    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="28" w:space="0" w:color="{border_color}"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
    tcPr.append(borders)

    p_title = cell.paragraphs[0]
    p_title.paragraph_format.line_spacing = 1.2
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(4)
    p_title.paragraph_format.first_line_indent = Cm(0)
    p_title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    r_title = p_title.add_run(f"{icon} {title}")
    r_title.bold = True
    r_title.font.name = 'Times New Roman'
    r_title.font.size = Pt(12)
    r_title.font.color.rgb = COLOR_ORANGE_HVNH if is_warning else COLOR_NAVY_HVNH

    p_body = cell.add_paragraph()
    p_body.paragraph_format.line_spacing = 1.25
    p_body.paragraph_format.space_before = Pt(0)
    p_body.paragraph_format.space_after = Pt(0)
    p_body.paragraph_format.first_line_indent = Cm(0)
    p_body.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    r_body = p_body.add_run(text)
    r_body.font.name = 'Times New Roman'
    r_body.font.size = Pt(12)
    r_body.font.color.rgb = COLOR_DARK_TEXT

    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_after = Pt(4)
    p_sp.paragraph_format.space_before = Pt(0)

# ==============================================================================
# CAPTION GẮN SỐ CHƯƠNG (Hình 3.4 / Bảng 2.1) — chuẩn luận văn HVNH
# ==============================================================================
def chapter_label(chapter_no, index_in_chapter):
    """Sinh nhãn đánh số gắn chương: ('3', 4) -> '3.4'. Dùng cho cả hình, bảng, phương trình."""
    return f"{chapter_no}.{index_in_chapter}"

def add_figure_with_caption(doc, img_path, fig_num, caption, source_text="(Nguồn: Kết quả phân tích từ tập dữ liệu thực nghiệm)", width=None):
    """Chèn hình ảnh kèm caption PHÍA DƯỚI hình + nguồn trích dẫn.
    fig_num nhận chuỗi gắn chương '3.4' (chuẩn luận văn HVNH) hoặc số nguyên (tương thích v1).
    """
    if width is None:
        width = Inches(5.9) if CURRENT_PROFILE is PROFILES["nd30-report"] else Cm(14.6)
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(10)
        p_img.paragraph_format.space_after = Pt(4)
        p_img.paragraph_format.first_line_indent = Cm(0)
        p_img.paragraph_format.keep_with_next = True
        run = p_img.add_run()
        run.add_picture(img_path, width=width)

    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(12)
    p_cap.paragraph_format.first_line_indent = Cm(0)

    r_num = p_cap.add_run(f"Hình {fig_num}: ")
    r_num.bold = True
    r_num.font.name = 'Times New Roman'
    r_num.font.size = Pt(11.5)
    r_num.font.color.rgb = COLOR_NAVY_HVNH

    r_cap = p_cap.add_run(caption + "\n")
    r_cap.font.name = 'Times New Roman'
    r_cap.font.size = Pt(11.5)
    r_cap.italic = True
    r_cap.font.color.rgb = COLOR_DARK_TEXT

    if source_text:
        r_src = p_cap.add_run(source_text)
        r_src.font.name = 'Times New Roman'
        r_src.font.size = Pt(10)
        r_src.italic = True
        r_src.font.color.rgb = COLOR_MUTED_TEXT

def _apply_table_body_style(table, zebra=True):
    """Định dạng thân bảng học thuật: tblHeader (lặp header qua trang), cantSplit,
    viền, padding, zebra striping, font ô bảng."""
    header_tr = table.rows[0]._tr.get_or_add_trPr()
    header_tr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

    for i, row in enumerate(table.rows):
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=90, bottom=90, left=110, right=110)
            tcPr = cell._tc.get_or_add_tcPr()
            borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/><w:bottom w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/><w:left w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/><w:right w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/></w:tcBorders>')
            tcPr.append(borders)

            for p in cell.paragraphs:
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.first_line_indent = Cm(0)
                if i == 0:
                    set_cell_background(cell, HEX_BLUE_HDR)
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs:
                        r.bold = True
                        r.font.name = 'Times New Roman'
                        r.font.size = Pt(10.5)
                        r.font.color.rgb = RGBColor(255, 255, 255)
                else:
                    if zebra and i % 2 == 1:
                        set_cell_background(cell, HEX_ZEBRA_BG)
                    for r in p.runs:
                        r.font.name = 'Times New Roman'
                        r.font.size = Pt(10.5)
                        r.font.color.rgb = COLOR_DARK_TEXT

def style_academic_table(doc, table, table_num, table_title, zebra=True):
    """Định dạng bảng học thuật kèm đầu đề PHÍA TRÊN bảng.
    table_num nhận chuỗi gắn chương '2.1' (chuẩn HVNH) hoặc số nguyên (tương thích v1)."""
    p_ttl = doc.add_paragraph()
    p_ttl.paragraph_format.space_before = Pt(10)
    p_ttl.paragraph_format.space_after = Pt(4)
    p_ttl.paragraph_format.keep_with_next = True
    p_ttl.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_ttl.paragraph_format.first_line_indent = Cm(0)

    r_tnum = p_ttl.add_run(f"Bảng {table_num}: ")
    r_tnum.bold = True
    r_tnum.font.name = 'Times New Roman'
    r_tnum.font.size = Pt(11.5)
    r_tnum.font.color.rgb = COLOR_NAVY_HVNH

    r_tname = p_ttl.add_run(table_title)
    r_tname.bold = True
    r_tname.font.name = 'Times New Roman'
    r_tname.font.size = Pt(11.5)
    r_tname.font.color.rgb = COLOR_DARK_TEXT

    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _apply_table_body_style(table, zebra=zebra)

    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_after = Pt(6)
    p_sp.paragraph_format.space_before = Pt(0)

# ==============================================================================
# TRANG ĐẦU: PHỤ BÌA, CAM ĐOAN, MỤC LỤC, DANH MỤC
# ==============================================================================
def add_title_page(doc, ten_de_tai, tac_gia, chuyen_nganh="", ma_so="",
                   nguoi_huong_dan=(), loai_luan_van="LUẬN VĂN THẠC SĨ KINH TẾ",
                   co_quan_chu_quan=("NGÂN HÀNG NHÀ NƯỚC VIỆT NAM", "BỘ GIÁO DỤC VÀ ĐÀO TẠO"),
                   noi_thang_ban="Hà Nội - năm 2026"):
    """Trang phụ bìa theo Mẫu 10 — Quy định trình bày luận văn HVNH.
    Tên đề tài IN HOA, không viết tắt. Bìa ngoài (đóng gáy in chữ nhũ) in ấn riêng theo Mẫu 09."""
    def _center(text, size=13, bold=False, space_after=0, color=COLOR_DARK_TEXT):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.first_line_indent = Cm(0)
        r = p.add_run(text.upper() if size >= 16 else text)
        r.bold = bold
        r.font.name = 'Times New Roman'
        r.font.size = Pt(size)
        r.font.color.rgb = color
        return p

    for i, cq in enumerate(co_quan_chu_quan):
        _center(cq, size=13, bold=True, space_after=0 if i < len(co_quan_chu_quan) - 1 else 6)
    _center("HỌC VIỆN NGÂN HÀNG", size=14, bold=True, space_after=72)
    _center(tac_gia, size=14, bold=True, space_after=48)
    _center(ten_de_tai, size=17, bold=True, space_after=36)
    if chuyen_nganh:
        _center(f"Chuyên ngành: {chuyen_nganh}", size=13, space_after=2)
    if ma_so:
        _center(f"Mã số: {ma_so}", size=13, space_after=24)
    _center(loai_luan_van, size=14, bold=True, space_after=36)
    if nguoi_huong_dan:
        _center("Người hướng dẫn khoa học:", size=13, bold=True, space_after=2)
        for i, gv in enumerate(nguoi_huong_dan):
            _center(f"{i + 1}. {gv}", size=13, space_after=2)
    _center("", space_after=60)
    _center(noi_thang_ban, size=13, bold=True)
    doc.add_page_break()

def add_cam_doan_page(doc, ho_ten, ten_de_tai, dia_diem="Hà Nội"):
    """Trang CAM ĐOAN danh dự — bắt buộc theo quy định chung của HVNH."""
    p_h = doc.add_paragraph("CAM ĐOAN")
    p_h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_h.paragraph_format.space_before = Pt(24)
    p_h.paragraph_format.space_after = Pt(18)
    p_h.paragraph_format.first_line_indent = Cm(0)
    for r in p_h.runs:
        r.bold = True
        r.font.name = 'Times New Roman'
        r.font.size = Pt(15)
        r.font.color.rgb = COLOR_NAVY_HVNH

    body = (
        f"Tôi cam đoan rằng công trình nghiên cứu \"{ten_de_tai}\" là sản phẩm do bản thân tôi "
        "trực tiếp thực hiện dưới sự hướng dẫn khoa học. Các số liệu và kết quả nêu trong luận văn "
        "là trung thực và chưa từng được công bố dưới bất kỳ hình thức nào. Nếu có sự vi phạm, "
        "tôi xin hoàn toàn chịu trách nhiệm trước Hội đồng khoa học và Học viện Ngân hàng."
    )
    add_body_paragraph(doc, body, indent=True)

    p_sign = doc.add_paragraph()
    p_sign.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_sign.paragraph_format.space_before = Pt(36)
    p_sig_loc = p_sign.add_run(f"{dia_diem}, ngày ...... tháng ...... năm ......")
    p_sig_loc.font.name = 'Times New Roman'
    p_sig_loc.font.size = Pt(13)
    p_sign2 = doc.add_paragraph()
    p_sign2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_sig_author = p_sign2.add_run("Tác giả luận văn\n\n\n")
    p_sig_author.font.name = 'Times New Roman'
    p_sig_author.font.size = Pt(13)
    r_name = p_sign2.add_run(ho_ten)
    r_name.bold = True
    r_name.font.name = 'Times New Roman'
    r_name.font.size = Pt(13)
    doc.add_page_break()

def add_toc_field(doc, heading_text="MỤC LỤC"):
    """Mục lục động đến 3 cấp số (TOC \\o '1-3') — chuẩn 'Mục lục đánh đến 3 số'.
    Word sẽ tự hỏi cập nhật khi mở file nhờ cờ updateFields; người dùng cũng có thể F9."""
    p_h = doc.add_paragraph(heading_text)
    p_h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_h.paragraph_format.space_after = Pt(18)
    p_h.paragraph_format.first_line_indent = Cm(0)
    for r in p_h.runs:
        r.bold = True
        r.font.name = 'Times New Roman'
        r.font.size = Pt(15)
        r.font.color.rgb = COLOR_NAVY_HVNH

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = CURRENT_PROFILE["line_spacing"]
    run = p.add_run()
    begin = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instr = parse_xml(r'<w:instrText %s xml:space="preserve"> TOC \o "1-3" \h \z \u </w:instrText>' % nsdecls('w'))
    sep = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    cached = parse_xml(r'<w:t %s>(Mở file trong Word và chọn Update Field / F9 để hiển thị mục lục)</w:t>' % nsdecls('w'))
    end = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    for el in (begin, instr, sep, cached, end):
        run._r.append(el)

    settings_el = doc.settings.element
    if not settings_el.findall(qn('w:updateFields')):
        settings_el.append(parse_xml('<w:updateFields %s w:val="true"/>' % nsdecls('w')))
    doc.add_page_break()

def add_abbreviation_table(doc, abbrev_items):
    """DANH MỤC CHỮ VIẾT TẮT xếp A→Z (quy định §2.4). abbrev_items: [(viết_tắt, nghĩa), ...]."""
    p_h = doc.add_paragraph("DANH MỤC CHỮ VIẾT TẮT")
    p_h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_h.paragraph_format.space_after = Pt(12)
    p_h.paragraph_format.first_line_indent = Cm(0)
    for r in p_h.runs:
        r.bold = True
        r.font.name = 'Times New Roman'
        r.font.size = Pt(15)
        r.font.color.rgb = COLOR_NAVY_HVNH

    items = sorted(abbrev_items, key=lambda kv: kv[0].upper())
    t = doc.add_table(rows=len(items) + 1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.rows[0].cells[0].paragraphs[0].add_run("Chữ viết tắt").bold = True
    t.rows[0].cells[1].paragraphs[0].add_run("Viết đầy đủ / Ý nghĩa").bold = True
    for i, (abbr, full) in enumerate(items, start=1):
        t.rows[i].cells[0].paragraphs[0].add_run(abbr)
        t.rows[i].cells[1].paragraphs[0].add_run(full)
    _apply_table_body_style(t, zebra=False)
    doc.add_page_break()

def add_exhibits_list(doc, heading_text, entries, label_prefix=None):
    """DANH MỤC BẢNG BIỂU / HÌNH VẼ dạng tĩnh. entries: [('2.1', 'Tiêu đề bảng/hình', trang), ...]."""
    p_h = doc.add_paragraph(heading_text)
    p_h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_h.paragraph_format.space_after = Pt(12)
    p_h.paragraph_format.first_line_indent = Cm(0)
    for r in p_h.runs:
        r.bold = True
        r.font.name = 'Times New Roman'
        r.font.size = Pt(15)
        r.font.color.rgb = COLOR_NAVY_HVNH

    for label, caption, page in entries:
        p = doc.add_paragraph()
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.left_indent = Cm(1.5)
        p.paragraph_format.first_line_indent = Cm(-1.5)
        p.paragraph_format.line_spacing = CURRENT_PROFILE["line_spacing"]
        p.paragraph_format.space_after = Pt(4)
        prefix = label_prefix or ("Bảng" if "BẢNG" in heading_text.upper() else "Hình")
        r_lbl = p.add_run(f"{prefix} {label}: ")
        r_lbl.bold = True
        r_lbl.font.name = 'Times New Roman'
        r_lbl.font.size = Pt(13)
        r_lbl.font.color.rgb = COLOR_NAVY_HVNH
        r_txt = p.add_run(caption)
        r_txt.font.name = 'Times New Roman'
        r_txt.font.size = Pt(13)
        r_txt.font.color.rgb = COLOR_DARK_TEXT
        if page:
            r_pg = p.add_run(f"  .......... {page}")
            r_pg.font.name = 'Times New Roman'
            r_pg.font.size = Pt(13)
    doc.add_page_break()

# ==============================================================================
# TÀI LIỆU THAM KHẢO CHUẨN HVNH (đánh số, tách ngôn ngữ, hanging indent 1cm)
# ==============================================================================
def hvnh_book(author, year, title, publisher, place):
    """Segment entry cho sách/luận án/báo cáo: Tác giả (Năm), Tên sách (nghiêng), NXB, Nơi XB."""
    return [
        (f"{author} ({year}), ", False),
        (title, True),
        (f", {publisher}, {place}.", False),
    ]

def hvnh_journal_article(author, year, article_title, journal, volume, issue, pages, english=False):
    """Segment entry cho bài báo tạp chí: Tác giả (Năm), \"Tên bài\" (ngoặc kép),
    Tên tạp chí (nghiêng), Tập(Số), tr./pp.trang."""
    page_prefix = "pp." if english else "tr."
    return [
        (f"{author} ({year}), ", False),
        (f"\u201c{article_title}\u201d", False),
        (", ", False),
        (journal, True),
        (f", {volume}({issue}), {page_prefix}{pages}.", False),
    ]

def add_reference_section_hvnh(doc, vietnamese_entries, foreign_entries):
    """Phần TÀI LIỆU THAM KHẢO chuẩn HVNH:
    - Đánh số liên tục xuyên hai nhóm ngôn ngữ (số này được dùng cho trích dẫn [n] trong thân bài)
    - Tách Tiếng Việt / Tiếng Anh; VN xếp ABC theo TÊN, nước ngoài theo HỌ (caller đã sort trước)
    - Entry dài hơn một dòng: dòng thứ hai lùi vào 1cm (hanging indent 1cm)
    vietnamese_entries / foreign_entries: danh sách segment trả về bởi hvnh_book / hvnh_journal_article
    """
    doc.add_page_break()
    p_h = doc.add_paragraph("TÀI LIỆU THAM KHẢO", style='Heading 1')
    p_h.paragraph_format.space_before = Pt(16)
    p_h.paragraph_format.first_line_indent = Cm(0)

    counter = {"n": 0}

    def _subhead(text):
        p_s = doc.add_paragraph(text)
        p_s.paragraph_format.space_before = Pt(12)
        p_s.paragraph_format.space_after = Pt(6)
        p_s.paragraph_format.first_line_indent = Cm(0)
        p_s.runs[0].bold = True
        p_s.runs[0].font.name = 'Times New Roman'
        p_s.runs[0].font.size = Pt(13.5)
        p_s.runs[0].font.color.rgb = COLOR_BLUE_SEC

    def _entry(segments):
        counter["n"] += 1
        p = doc.add_paragraph()
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.left_indent = Cm(1.0)
        p.paragraph_format.first_line_indent = Cm(-1.0)
        p.paragraph_format.line_spacing = CURRENT_PROFILE["line_spacing"]
        p.paragraph_format.space_after = Pt(6)
        r_num = p.add_run(f"{counter['n']}. ")
        r_num.font.name = 'Times New Roman'
        r_num.font.size = Pt(13)
        r_num.font.color.rgb = COLOR_DARK_TEXT
        for txt, italic in segments:
            r = p.add_run(txt)
            r.italic = italic
            r.font.name = 'Times New Roman'
            r.font.size = Pt(13)
            r.font.color.rgb = COLOR_DARK_TEXT
        return p

    _subhead("Tiếng Việt")
    for seg in vietnamese_entries:
        _entry(seg)
    if foreign_entries:
        _subhead("Tiếng Anh")
        for seg in foreign_entries:
            _entry(seg)

def add_published_works_section(doc, works):
    """DANH MỤC CÔNG TRÌNH KHOA HỌC CỦA TÁC GIẢ ĐÃ CÔNG BỐ — đứng trước TLTK theo trình tự."""
    doc.add_page_break()
    p_h = doc.add_paragraph("DANH MỤC CÔNG TRÌNH KHOA HỌC CỦA TÁC GIẢ ĐÃ CÔNG BỐ", style='Heading 1')
    p_h.paragraph_format.first_line_indent = Cm(0)
    for i, w in enumerate(works, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.left_indent = Cm(1.0)
        p.paragraph_format.first_line_indent = Cm(-1.0)
        p.paragraph_format.space_after = Pt(4)
        r_n = p.add_run(f"{i}. ")
        r_n.font.name = 'Times New Roman'
        r_n.font.size = Pt(13)
        r_w = p.add_run(w)
        r_w.font.name = 'Times New Roman'
        r_w.font.size = Pt(13)
