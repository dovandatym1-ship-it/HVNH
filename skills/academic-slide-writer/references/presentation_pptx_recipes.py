"""
Antigravity Academic Writer - Python-pptx Modular Recipe Library for Executive Academic Presentations
Implements 16:9 Widescreen, 4-Layer Architecture, Action Titles, Card Grid & Speaker Notes
v2 — thêm đủ 6 Content Variants (Model Comparison, Strategic Tier Matrix), cảnh báo rõ ràng
     khi thiếu ảnh / lệch số lượng thẻ thiết kế, và validate_deck() để tự kiểm tra deck
     trước khi giao (kích thước 16:9, Speaker Notes đủ độ dài) — xem SKILL.md §2.2 và §2.4.
"""
import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

# ==============================================================================
# HỆ MÀU NHẬN DIỆN THƯƠNG HIỆU HỌC VIỆN NGÂN HÀNG & EXECUTIVE DASHBOARD
# ==============================================================================
C_HVNH_NAVY    = RGBColor(10, 37, 64)       # #0A2540: Xanh Navy đậm BAV
C_HVNH_ORANGE  = RGBColor(234, 88, 12)      # #EA580C: Cam năng lượng BAV
C_HVNH_GOLD    = RGBColor(245, 158, 11)     # #F59E0B: Vàng điểm nhấn
C_BLUE_CYAN    = RGBColor(2, 132, 199)      # #0284C7: Xanh Cyan chỉ số phụ
C_ACCENT_RED   = RGBColor(220, 38, 38)      # #DC2626: Đỏ cảnh báo rủi ro
C_ACCENT_GREEN = RGBColor(16, 185, 129)     # #10B981: Xanh lá thành công

C_CANVAS_BG    = RGBColor(248, 250, 252)    # #F8FAFC: Nền canvas Slate 50
C_CARD_BG      = RGBColor(255, 255, 255)    # #FFFFFF: Nền thẻ Card trắng
C_CARD_BORDER  = RGBColor(226, 232, 240)    # #E2E8F0: Viền thẻ Slate 200
C_TEXT_DARK    = RGBColor(15, 23, 42)       # #0F172A: Màu chữ Slate 900
C_TEXT_MUTED   = RGBColor(100, 116, 139)    # #64748B: Màu chữ chú thích Slate 500
C_WHITE        = RGBColor(255, 255, 255)

FONT_PRIMARY = 'Calibri'
FONT_TITLE   = 'Calibri'

def create_executive_deck():
    """Khởi tạo Deck chuẩn Widescreen 16:9 (13.333 x 7.5 inches)"""
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs

def add_blank_slide(prs, bg_color=C_CANVAS_BG):
    """Tạo slide trống với màu nền canvas tinh chỉnh"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = bg_color
    bg.line.fill.background()
    return slide

def add_action_header(slide, tag_text, title_text, page_num=None, total_pages=20):
    """
    Tạo Header chuẩn Action Title của EasySlides / McKinsey:
    - Pill Badge ở trên (Tag)
    - Action Title đậm nét đúc kết thông điệp chính ở dưới
    - Page Number đếm trang ở góc phải
    """
    # Pill Badge Tag
    tag_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.42), Inches(3.2), Inches(0.32))
    tag_box.fill.solid()
    tag_box.fill.fore_color.rgb = C_HVNH_NAVY
    tag_box.line.fill.background()
    tf_tag = tag_box.text_frame
    tf_tag.word_wrap = True
    tf_tag.margin_left = Inches(0.1)
    tf_tag.margin_top = Inches(0.04)
    p_tag = tf_tag.paragraphs[0]
    p_tag.alignment = PP_ALIGN.CENTER
    r_tag = p_tag.add_run()
    r_tag.text = tag_text.upper()
    r_tag.font.name = FONT_PRIMARY
    r_tag.font.size = Pt(9.5)
    r_tag.font.bold = True
    r_tag.font.color.rgb = C_WHITE

    # Action Title Text
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.78), Inches(10.5), Inches(0.75))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0)
    tf.margin_top = Inches(0)
    p_title = tf.paragraphs[0]
    p_title.alignment = PP_ALIGN.LEFT
    r_title = p_title.add_run()
    r_title.text = title_text
    r_title.font.name = FONT_TITLE
    r_title.font.size = Pt(20)
    r_title.font.bold = True
    r_title.font.color.rgb = C_HVNH_NAVY

    # Page Number
    if page_num:
        tb_p = slide.shapes.add_textbox(Inches(11.5), Inches(0.42), Inches(1.2), Inches(0.35))
        p_p = tb_p.text_frame.paragraphs[0]
        p_p.alignment = PP_ALIGN.RIGHT
        r_p = p_p.add_run()
        r_p.text = f"{page_num:02d} / {total_pages:02d}"
        r_p.font.name = FONT_PRIMARY
        r_p.font.size = Pt(11)
        r_p.font.bold = True
        r_p.font.color.rgb = C_TEXT_MUTED

def add_card_box(slide, left, top, width, height, bg_color=C_CARD_BG, border_color=C_CARD_BORDER, border_width=Pt(1.2)):
    """Tạo khung Card bo góc viền tinh tế"""
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    if border_color:
        card.line.color.rgb = border_color
        card.line.width = border_width
    else:
        card.line.fill.background()
    return card

def add_metric_card(slide, left, top, width, height, lead_kpi, kpi_label, desc_lines, highlight_color=C_HVNH_ORANGE):
    """Tạo thẻ Metric với số liệu lớn (Lead KPI) nổi bật"""
    add_card_box(slide, left, top, width, height)

    # Top Accent Strip
    strip = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left + Inches(0.15), top + Inches(0.15), width - Inches(0.3), Inches(0.06))
    strip.fill.solid()
    strip.fill.fore_color.rgb = highlight_color
    strip.line.fill.background()

    tb = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.3), width - Inches(0.4), height - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True

    # KPI Number
    p_kpi = tf.paragraphs[0]
    p_kpi.alignment = PP_ALIGN.LEFT
    r_kpi = p_kpi.add_run()
    r_kpi.text = lead_kpi
    r_kpi.font.name = FONT_TITLE
    r_kpi.font.size = Pt(30)
    r_kpi.font.bold = True
    r_kpi.font.color.rgb = highlight_color

    # KPI Label
    p_lbl = tf.add_paragraph()
    p_lbl.space_before = Pt(2)
    p_lbl.space_after = Pt(6)
    r_lbl = p_lbl.add_run()
    r_lbl.text = kpi_label
    r_lbl.font.name = FONT_PRIMARY
    r_lbl.font.size = Pt(12)
    r_lbl.font.bold = True
    r_lbl.font.color.rgb = C_HVNH_NAVY

    # Descriptions
    for line in desc_lines:
        p_d = tf.add_paragraph()
        p_d.space_before = Pt(2)
        r_d = p_d.add_run()
        r_d.text = f"• {line}"
        r_d.font.name = FONT_PRIMARY
        r_d.font.size = Pt(10.5)
        r_d.font.color.rgb = C_TEXT_DARK

def add_chart_deep_dive_60_40(slide, img_path, left_notes_title, note_items):
    """Bố cục Chart Deep-Dive 60/40: 60% biểu đồ hình ảnh bên trái, 40% phân tích bên phải"""
    # 60% Image on Left
    if os.path.exists(img_path):
        add_card_box(slide, Inches(0.8), Inches(1.65), Inches(7.4), Inches(5.2))
        slide.shapes.add_picture(img_path, Inches(0.95), Inches(1.8), width=Inches(7.1))
    else:
        print(f"[WARN] add_chart_deep_dive_60_40: không tìm thấy ảnh tại '{img_path}' — bỏ qua khung ảnh 60%, slide sẽ thiếu biểu đồ chính.")

    # 40% Card on Right
    add_card_box(slide, Inches(8.4), Inches(1.65), Inches(4.1), Inches(5.2))
    tb = slide.shapes.add_textbox(Inches(8.6), Inches(1.85), Inches(3.7), Inches(4.8))
    tf = tb.text_frame
    tf.word_wrap = True

    p_hdr = tf.paragraphs[0]
    r_hdr = p_hdr.add_run()
    r_hdr.text = left_notes_title
    r_hdr.font.name = FONT_TITLE
    r_hdr.font.size = Pt(14)
    r_hdr.font.bold = True
    r_hdr.font.color.rgb = C_HVNH_NAVY

    for tag, desc in note_items:
        p_item = tf.add_paragraph()
        p_item.space_before = Pt(8)

        r_tag = p_item.add_run()
        r_tag.text = f"📌 {tag}: "
        r_tag.font.name = FONT_PRIMARY
        r_tag.font.size = Pt(11)
        r_tag.font.bold = True
        r_tag.font.color.rgb = C_HVNH_ORANGE

        r_desc = p_item.add_run()
        r_desc.text = desc
        r_desc.font.name = FONT_PRIMARY
        r_desc.font.size = Pt(10.5)
        r_desc.font.color.rgb = C_TEXT_DARK

def add_process_pipeline_4(slide, steps_data):
    """
    Tạo quy trình 4 bước ngang (Horizontal 4-Step Pipeline)
    steps_data = [("01", "TIỀN XỬ LÝ", ["Loại bỏ 2 cột rò rỉ", "Điền khuyết thiếu"]), ...]
    Hàm được thiết kế cố định cho đúng 4 bước (bố cục 4 cột chia đều khung 16:9);
    số bước khác 4 vẫn chạy được nhưng có thể lệch khung, sẽ in cảnh báo.
    """
    if len(steps_data) != 4:
        print(f"[WARN] add_process_pipeline_4: nhận {len(steps_data)} bước, hàm được thiết kế cho đúng 4 bước — bố cục có thể tràn hoặc hụt khung slide nếu khác 4.")

    card_w = Inches(2.75)
    card_h = Inches(5.1)
    top_pos = Inches(1.7)

    for i, (num_str, title, lines) in enumerate(steps_data):
        left_pos = Inches(0.8) + i * Inches(2.95)
        add_card_box(slide, left_pos, top_pos, card_w, card_h)

        # Step Badge
        badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_pos + Inches(0.18), top_pos + Inches(0.18), Inches(0.7), Inches(0.4))
        badge.fill.solid()
        badge.fill.fore_color.rgb = C_HVNH_ORANGE if i == 0 else C_HVNH_NAVY
        badge.line.fill.background()
        tf_b = badge.text_frame
        p_b = tf_b.paragraphs[0]
        p_b.alignment = PP_ALIGN.CENTER
        r_b = p_b.add_run()
        r_b.text = num_str
        r_b.font.bold = True
        r_b.font.color.rgb = C_WHITE

        # Text Content
        tb = slide.shapes.add_textbox(left_pos + Inches(0.18), top_pos + Inches(0.7), card_w - Inches(0.36), card_h - Inches(0.9))
        tf = tb.text_frame
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        r_t = p_t.add_run()
        r_t.text = title
        r_t.font.bold = True
        r_t.font.size = Pt(13)
        r_t.font.color.rgb = C_HVNH_NAVY

        for line in lines:
            p_l = tf.add_paragraph()
            p_l.space_before = Pt(6)
            r_l = p_l.add_run()
            r_l.text = f"• {line}"
            r_l.font.size = Pt(10.5)
            r_l.font.color.rgb = C_TEXT_DARK

def add_persona_profile_cards(slide, personas_data):
    """
    Tạo 3 thẻ Chân dung khách hàng Personas
    personas_data = [("1. Silent Defector", "55% Rời bỏ", "Dư nợ $0, quẹt thẻ <45 lần", "E-voucher cuối tuần, hoàn tiền 5%"), ...]
    Hàm được thiết kế cố định cho đúng 3 thẻ (bố cục chia đều khung 16:9);
    số thẻ khác 3 vẫn chạy được nhưng có thể lệch khung, sẽ in cảnh báo.
    """
    if len(personas_data) != 3:
        print(f"[WARN] add_persona_profile_cards: nhận {len(personas_data)} persona, hàm được thiết kế cho đúng 3 thẻ — bố cục có thể tràn hoặc hụt khung slide nếu khác 3.")

    card_w = Inches(3.7)
    card_h = Inches(5.1)
    top_pos = Inches(1.7)

    for i, (name, share, features, action) in enumerate(personas_data):
        left_pos = Inches(0.8) + i * Inches(3.95)
        add_card_box(slide, left_pos, top_pos, card_w, card_h)

        # Header Box in Card
        hdr_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_pos + Inches(0.15), top_pos + Inches(0.15), card_w - Inches(0.3), Inches(0.85))
        hdr_box.fill.solid()
        hdr_box.fill.fore_color.rgb = C_HVNH_NAVY if i != 0 else C_HVNH_ORANGE
        hdr_box.line.fill.background()
        tf_h = hdr_box.text_frame
        p_h0 = tf_h.paragraphs[0]
        p_h0.alignment = PP_ALIGN.CENTER
        r_h0 = p_h0.add_run()
        r_h0.text = name
        r_h0.font.bold = True
        r_h0.font.size = Pt(12)
        r_h0.font.color.rgb = C_WHITE

        p_h1 = tf_h.add_paragraph()
        p_h1.alignment = PP_ALIGN.CENTER
        r_h1 = p_h1.add_run()
        r_h1.text = share
        r_h1.font.bold = True
        r_h1.font.size = Pt(11)
        r_h1.font.color.rgb = C_WHITE

        # Details
        tb = slide.shapes.add_textbox(left_pos + Inches(0.2), top_pos + Inches(1.15), card_w - Inches(0.4), card_h - Inches(1.3))
        tf = tb.text_frame
        tf.word_wrap = True

        p_f_lbl = tf.paragraphs[0]
        r_f_lbl = p_f_lbl.add_run()
        r_f_lbl.text = "🔍 Đặc trưng hành vi:"
        r_f_lbl.font.bold = True
        r_f_lbl.font.size = Pt(11)
        r_f_lbl.font.color.rgb = C_HVNH_NAVY

        p_f_val = tf.add_paragraph()
        p_f_val.space_before = Pt(2)
        p_f_val.space_after = Pt(10)
        r_f_val = p_f_val.add_run()
        r_f_val.text = features
        r_f_val.font.size = Pt(10.5)
        r_f_val.font.color.rgb = C_TEXT_DARK

        p_a_lbl = tf.add_paragraph()
        r_a_lbl = p_a_lbl.add_run()
        r_a_lbl.text = "🎯 Hành động can thiệp:"
        r_a_lbl.font.bold = True
        r_a_lbl.font.size = Pt(11)
        r_a_lbl.font.color.rgb = C_HVNH_ORANGE

        p_a_val = tf.add_paragraph()
        p_a_val.space_before = Pt(2)
        r_a_val = p_a_val.add_run()
        r_a_val.text = action
        r_a_val.font.size = Pt(10.5)
        r_a_val.font.color.rgb = C_TEXT_DARK

def add_model_comparison_cards(slide, models_data, winner_index=0):
    """
    Variant 4 — Side-by-Side Model Comparison: so sánh N mô hình cạnh nhau,
    thẻ Winner được làm nổi bật viền Cam HVNH + huy hiệu "★ WINNER".
    models_data = [
        ("XGBoost", [("ROC-AUC", "0.9930"), ("Recall", "91.38%"), ("F1-Score", "0.887")]),
        ("Random Forest", [("ROC-AUC", "0.9711"), ("Recall", "85.20%"), ("F1-Score", "0.841")]),
        ...
    ]
    winner_index: chỉ số (0-based) của mô hình được đánh dấu Winner. Mặc định 0.
    Bố cục chia đều khung 16:9 theo số mô hình truyền vào (khuyến nghị 2-4 mô hình).
    """
    n = len(models_data)
    if n == 0:
        print("[WARN] add_model_comparison_cards: models_data rỗng — không có thẻ nào được vẽ.")
        return
    if n > 5:
        print(f"[WARN] add_model_comparison_cards: {n} mô hình có thể khiến thẻ quá hẹp để đọc (khuyến nghị 2-4).")
    if not (0 <= winner_index < n):
        print(f"[WARN] add_model_comparison_cards: winner_index={winner_index} nằm ngoài phạm vi hợp lệ [0, {n - 1}] — không thẻ nào được đánh dấu Winner.")
        winner_index = -1

    gap_in = 0.2
    total_gap_in = gap_in * (n - 1) if n > 1 else 0
    card_w = Inches((13.333 - 1.6 - total_gap_in) / n)
    card_h = Inches(5.1)
    top_pos = Inches(1.7)

    for i, (model_name, metrics) in enumerate(models_data):
        left_pos = Inches(0.8 + i * (card_w.inches + gap_in))
        is_winner = (i == winner_index)
        border_color = C_HVNH_ORANGE if is_winner else C_CARD_BORDER
        border_width = Pt(2.4) if is_winner else Pt(1.2)
        add_card_box(slide, left_pos, top_pos, card_w, card_h, border_color=border_color, border_width=border_width)

        if is_winner:
            badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_pos + Inches(0.15), top_pos - Inches(0.02), card_w - Inches(0.3), Inches(0.34))
            badge.fill.solid()
            badge.fill.fore_color.rgb = C_HVNH_ORANGE
            badge.line.fill.background()
            tf_b = badge.text_frame
            p_b = tf_b.paragraphs[0]
            p_b.alignment = PP_ALIGN.CENTER
            r_b = p_b.add_run()
            r_b.text = "★ WINNER"
            r_b.font.name = FONT_PRIMARY
            r_b.font.size = Pt(10)
            r_b.font.bold = True
            r_b.font.color.rgb = C_WHITE

        tb = slide.shapes.add_textbox(left_pos + Inches(0.2), top_pos + Inches(0.45), card_w - Inches(0.4), card_h - Inches(0.65))
        tf = tb.text_frame
        tf.word_wrap = True

        p_name = tf.paragraphs[0]
        r_name = p_name.add_run()
        r_name.text = model_name
        r_name.font.name = FONT_TITLE
        r_name.font.size = Pt(15)
        r_name.font.bold = True
        r_name.font.color.rgb = C_HVNH_ORANGE if is_winner else C_HVNH_NAVY

        for metric_label, metric_value in metrics:
            p_val = tf.add_paragraph()
            p_val.space_before = Pt(10)
            r_val = p_val.add_run()
            r_val.text = metric_value
            r_val.font.name = FONT_TITLE
            r_val.font.size = Pt(18)
            r_val.font.bold = True
            r_val.font.color.rgb = C_TEXT_DARK

            p_lbl = tf.add_paragraph()
            p_lbl.space_before = Pt(0)
            r_lbl = p_lbl.add_run()
            r_lbl.text = metric_label
            r_lbl.font.name = FONT_PRIMARY
            r_lbl.font.size = Pt(10)
            r_lbl.font.color.rgb = C_TEXT_MUTED

def add_strategic_tier_matrix(slide, tiers_data):
    """
    Variant 6 — Strategic Tier Matrix: khung giải pháp xếp tầng theo chiều dọc,
    mỗi tầng gồm nhãn tầng (trái), mô tả hành động (giữa) và chỉ số KPI đo lường (phải).
    tiers_data = [
        ("Tầng 1: EWS", "Hệ thống cảnh báo sớm tự động chấm điểm rủi ro hằng ngày", "Độ chính xác > 90%"),
        ("Tầng 2: CSKH Chủ động", "Gọi điện chăm sóc khi chạm ngưỡng rủi ro trung bình", "Giảm 20% khiếu nại"),
        ...
    ]
    Khuyến nghị tối đa 4-5 tầng để vừa khung slide 16:9; màu tầng luân phiên theo bảng màu HVNH.
    """
    n = len(tiers_data)
    if n == 0:
        print("[WARN] add_strategic_tier_matrix: tiers_data rỗng — không có tầng nào được vẽ.")
        return
    if n > 5:
        print(f"[WARN] add_strategic_tier_matrix: {n} tầng có thể tràn khung slide theo chiều dọc (khuyến nghị ≤ 5).")

    top_start = Inches(1.65)
    total_h = Inches(5.2)
    gap = Inches(0.12)
    tier_h = Inches((5.2 - 0.12 * (n - 1)) / n)
    tier_colors = [C_HVNH_NAVY, C_BLUE_CYAN, C_HVNH_GOLD, C_HVNH_ORANGE, C_ACCENT_GREEN]

    for i, (tier_title, tier_desc, tier_kpi) in enumerate(tiers_data):
        top_pos = Inches(top_start.inches + i * (tier_h.inches + gap.inches))
        color = tier_colors[i % len(tier_colors)]

        label_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), top_pos, Inches(2.6), tier_h)
        label_box.fill.solid()
        label_box.fill.fore_color.rgb = color
        label_box.line.fill.background()
        tf_l = label_box.text_frame
        tf_l.word_wrap = True
        p_l = tf_l.paragraphs[0]
        p_l.alignment = PP_ALIGN.CENTER
        r_l = p_l.add_run()
        r_l.text = tier_title
        r_l.font.name = FONT_TITLE
        r_l.font.size = Pt(12.5)
        r_l.font.bold = True
        r_l.font.color.rgb = C_WHITE

        add_card_box(slide, Inches(3.55), top_pos, Inches(6.35), tier_h, border_color=color, border_width=Pt(1.5))
        tb_d = slide.shapes.add_textbox(Inches(3.75), top_pos + Inches(0.08), Inches(5.95), Inches(max(tier_h.inches - 0.16, 0.2)))
        tf_d = tb_d.text_frame
        tf_d.word_wrap = True
        p_d = tf_d.paragraphs[0]
        r_d = p_d.add_run()
        r_d.text = tier_desc
        r_d.font.name = FONT_PRIMARY
        r_d.font.size = Pt(11)
        r_d.font.color.rgb = C_TEXT_DARK

        add_card_box(slide, Inches(10.1), top_pos, Inches(2.43), tier_h, bg_color=C_CANVAS_BG, border_color=color)
        tb_k = slide.shapes.add_textbox(Inches(10.2), top_pos + Inches(0.05), Inches(2.23), Inches(max(tier_h.inches - 0.1, 0.2)))
        tf_k = tb_k.text_frame
        tf_k.word_wrap = True
        p_k = tf_k.paragraphs[0]
        p_k.alignment = PP_ALIGN.CENTER
        r_k = p_k.add_run()
        r_k.text = tier_kpi
        r_k.font.name = FONT_PRIMARY
        r_k.font.size = Pt(10.5)
        r_k.font.bold = True
        r_k.font.color.rgb = color

def add_speaker_notes(slide, notes_text):
    """Chèn lời thoại thuyết trình (Speaker Notes) tích hợp sẵn trong Slide"""
    notes_slide = slide.notes_slide
    text_frame = notes_slide.notes_text_frame
    text_frame.text = notes_text

# ==============================================================================
# TỰ KIỂM TRA DECK TRƯỚC KHI GIAO (QA Gate)
# ==============================================================================
def validate_deck(prs, expected_slides=None, min_notes_words=40, max_notes_words=170):
    """
    Tự dò lại deck vừa dựng và in checklist Đạt/Không đạt trước khi xuất file .pptx:
      - Kích thước slide đúng chuẩn Widescreen 16:9 (13.333 x 7.5 inch).
      - Số slide khớp với kỳ vọng (nếu truyền expected_slides, ví dụ 20 theo Master Blueprint).
      - Mỗi slide có Speaker Notes, độ dài nằm trong khoảng khuyến nghị 80-150 từ
        (min_notes_words/max_notes_words nới lỏng hơn một chút để không quá khắt khe).
    Trả về True nếu không phát hiện vấn đề, False nếu có — luôn in báo cáo ra console.
    Lưu ý: hàm KHÔNG thể kiểm tra nội dung Action Title có "đúng theo Ghost Deck Test" hay
    không — việc đó cần con người đọc lại chuỗi tiêu đề (xem SKILL.md §2.1).
    """
    issues = []

    is_widescreen = prs.slide_width == Inches(13.333) and prs.slide_height == Inches(7.5)
    if not is_widescreen:
        issues.append(
            f"Kích thước slide không phải 16:9 chuẩn (13.333x7.5in): hiện tại "
            f"{prs.slide_width.inches:.3f}x{prs.slide_height.inches:.3f} inch."
        )

    n_slides = len(prs.slides)
    if expected_slides and n_slides != expected_slides:
        issues.append(f"Số slide thực tế ({n_slides}) khác số dự kiến ({expected_slides}).")

    for idx, slide in enumerate(prs.slides, start=1):
        notes = ""
        if slide.has_notes_slide:
            notes = (slide.notes_slide.notes_text_frame.text or "").strip()
        word_count = len(notes.split()) if notes else 0
        if word_count == 0:
            issues.append(f"Slide {idx}: chưa có Speaker Notes.")
        elif word_count < min_notes_words:
            issues.append(f"Slide {idx}: Speaker Notes chỉ {word_count} từ (khuyến nghị {min_notes_words}-{max_notes_words} từ).")
        elif word_count > max_notes_words:
            issues.append(f"Slide {idx}: Speaker Notes dài {word_count} từ, vượt khuyến nghị {min_notes_words}-{max_notes_words} từ.")

    print("=== VALIDATE DECK: BÁO CÁO KIỂM TRA ===")
    print(f"Tổng số slide: {n_slides}")
    if not issues:
        print("✅ ĐẠT — không phát hiện lỗi cấu trúc.")
    else:
        print(f"⚠️ PHÁT HIỆN {len(issues)} VẤN ĐỀ:")
        for iss in issues:
            print(f"  - {iss}")
    return len(issues) == 0
